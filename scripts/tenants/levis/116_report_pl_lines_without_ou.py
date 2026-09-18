# -*- coding: utf-8 -*-
"""Baris P&L terposting yang tidak punya Operating Unit.

Sheet baris #36. SELECT-only -- tidak pernah menulis apa pun.

  docker exec -i -e OUT=/mnt/data_levis/PL_Tanpa_Operating_Unit.xlsx \\
    odoo19-platform-odoo odoo shell -d prd_levis_begbal \\
    --no-http --max-cron-threads=0 --shell-interface=python \\
    < scripts/tenants/levis/116_report_pl_lines_without_ou.py

Dijalankan lewat ``odoo shell`` (butuh ``env``), bukan python biasa.

--------------------------------------------------------------------------
Untuk apa
--------------------------------------------------------------------------
Saklar ``custom_levis_localization.ou_required_pl`` menolak posting baris P&L
tanpa Operating Unit. Sebelum saklar itu dinyalakan, Finance perlu tahu baris
mana yang hari ini sudah terlanjur begitu -- bukan karena baris lama akan
ditolak (posting sudah terjadi), melainkan karena baris-baris itulah yang
membuat laporan per toko tidak lengkap, dan pola jurnalnya menunjukkan jalur
mana yang masih membiarkan OU kosong.

Per 18-Sep-2026 di ``prd_levis_begbal``: 21.988 baris P&L terposting 2026,
**79** tanpa OU -- GLJV 50, OBCA 12, EBRTB 6, IBRI/IMand/IBNI/BILL 9, DEPRE 2.
Yang dari jurnal bank itulah "bank admin" yang dilaporkan klien 16-Sep.

Kolom akun dibaca per company (``code`` company-dependent di Odoo 19).
"""

import os

OUT = os.environ.get("OUT", "/mnt/data_levis/PL_Tanpa_Operating_Unit.xlsx")
DATE_FROM = os.environ.get("DATE_FROM", "2026-01-01")

Move = env["account.move"]
company = env["res.company"].search([], order="id", limit=1)
ou_ids = Move._levis_ou_analytic_ids()
pl_ids = Move._levis_pl_account_ids(company)

print("=" * 72)
print("Baris P&L tanpa Operating Unit  --  sejak %s" % DATE_FROM)
print("=" * 72)
print("Company: %s" % company.display_name)
print("Akun P&L (kode 5-9): %s" % len(pl_ids))
print("Akun analytic di plan Operating Unit: %s" % len(ou_ids))

lines = env["account.move.line"].search(
    [
        ("company_id", "=", company.id),
        ("account_id", "in", list(pl_ids)),
        ("parent_state", "=", "posted"),
        ("date", ">=", DATE_FROM),
        ("display_type", "not in", ("line_section", "line_note")),
    ]
)
print("\nBaris P&L terposting: %s" % len(lines))

missing = lines.filtered(
    lambda ln: not ln.l10n_ou_analytic_id and not Move._levis_distribution_has_ou(ln.analytic_distribution, ou_ids)
)
print("Tanpa Operating Unit: %s" % len(missing))

per_journal = {}
for line in missing:
    code = line.move_id.journal_id.code
    bucket = per_journal.setdefault(code, {"lines": 0, "amount": 0.0})
    bucket["lines"] += 1
    bucket["amount"] += abs(line.balance)

print("\nPer jurnal:")
for code in sorted(per_journal, key=lambda c: -per_journal[c]["lines"]):
    bucket = per_journal[code]
    print("  %-8s %5s baris  %s" % (code, bucket["lines"], company.currency_id.round(bucket["amount"])))

per_account = {}
for line in missing:
    code = line.account_id.with_company(company).code
    bucket = per_account.setdefault(code, {"lines": 0, "amount": 0.0, "name": line.account_id.display_name})
    bucket["lines"] += 1
    bucket["amount"] += abs(line.balance)

print("\nPer akun:")
for code in sorted(per_account, key=lambda c: -per_account[c]["lines"]):
    bucket = per_account[code]
    print("  %-14s %4s baris  %s" % (code, bucket["lines"], company.currency_id.round(bucket["amount"])))

try:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Baris"
    ws.append(["Tanggal", "Jurnal", "Nomor", "Akun", "Nama Akun", "Label", "Debit", "Kredit", "Partner"])
    for line in missing.sorted(lambda ln: (ln.date, ln.move_id.name or "")):
        ws.append(
            [
                line.date,
                line.move_id.journal_id.code,
                line.move_id.name,
                line.account_id.with_company(company).code,
                line.account_id.display_name,
                line.name or "",
                line.debit,
                line.credit,
                line.partner_id.display_name or "",
            ]
        )
    w2 = wb.create_sheet("Per Jurnal")
    w2.append(["Jurnal", "Baris", "Nilai absolut"])
    for code in sorted(per_journal, key=lambda c: -per_journal[c]["lines"]):
        w2.append([code, per_journal[code]["lines"], per_journal[code]["amount"]])
    w3 = wb.create_sheet("Per Akun")
    w3.append(["Akun", "Nama", "Baris", "Nilai absolut"])
    for code in sorted(per_account, key=lambda c: -per_account[c]["lines"]):
        w3.append([code, per_account[code]["name"], per_account[code]["lines"], per_account[code]["amount"]])
    wb.save(OUT)
    print("\nLaporan: %s" % OUT)
except Exception as exc:  # noqa: BLE001
    print("\n(laporan xlsx gagal: %s)" % str(exc)[:70])

env.cr.rollback()
