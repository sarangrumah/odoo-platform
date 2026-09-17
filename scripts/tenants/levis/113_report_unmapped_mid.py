# -*- coding: utf-8 -*-
"""MID bank yang belum terpetakan ke toko -- SELECT-ONLY.

Sheet baris #41. Satu baris per MID, bukan per mutasi: memetakan satu MID
menyelesaikan **setiap** baris yang membawanya, lalu dan nanti, jadi daftar
per-MID adalah daftar kerja yang sebenarnya. Kolom "Contoh Narasi" ada supaya
Finance AR bisa mengenali tokonya tanpa membuka Odoo.

  python3 scripts/tenants/levis/113_report_unmapped_mid.py

Env:  DB   -> database (default prd_levis_begbal)
      OUT  -> path xlsx

Tidak menulis apa pun ke Odoo; satu-satunya tulisan adalah berkas Excel di
/srv/sftp-share/files (bisa diunduh lewat File Browser /files).
"""

import csv
import io
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PG = "odoo19-platform-postgres"
DB = os.environ.get("DB", "prd_levis_begbal")
OUT = os.environ.get("OUT", "/srv/sftp-share/files/MID_Belum_Terpetakan.xlsx")

HDR = PatternFill("solid", fgColor="1F4E78")
HF = Font(bold=True, color="FFFFFF")
WARN = PatternFill("solid", fgColor="FCE4D6")
NOTE = Font(italic=True, color="595959")


def q(sql):
    out = subprocess.run(
        [
            "docker",
            "exec",
            PG,
            "sh",
            "-c",
            f'PGPASSWORD="$POSTGRES_PASSWORD" psql -U "$POSTGRES_USER" -d {DB} --csv -v ON_ERROR_STOP=1 -c "{sql}"',
        ],
        capture_output=True,
        text=True,
    )
    if out.returncode:
        sys.exit("query gagal:\n" + out.stderr)
    return list(csv.DictReader(io.StringIO(out.stdout)))


# Satu baris per MID. ``levis_gross`` adalah nilai bruto sebelum MDR bila
# narasinya menyebutkannya; kalau tidak, amount statement yang dipakai.
UNMAPPED = """
select l.levis_mid as mid,
       count(*) as baris,
       min(m.date)::text as pertama,
       max(m.date)::text as terakhir,
       round(sum(coalesce(l.levis_gross, l.amount))::numeric, 2) as nilai,
       min(l.payment_ref) as contoh_narasi,
       coalesce(min(j.name->>'en_US'), min(j.code), '') as jurnal
from account_bank_statement_line l
join account_move m on m.id = l.move_id
left join account_journal j on j.id = l.journal_id
where l.levis_mid is not null and l.levis_mid <> ''
  and l.levis_ou_analytic_id is null
group by l.levis_mid
order by 5 desc
"""

TOTALS = """
select count(*) as baris,
       count(*) filter (where levis_mid is not null and levis_mid <> '') as ber_mid,
       count(*) filter (where levis_mid is not null and levis_mid <> ''
                          and levis_ou_analytic_id is null) as mid_tanpa_ou,
       count(*) filter (where levis_ou_analytic_id is null
                          and (levis_mid is null or levis_mid = '')) as tanpa_mid_tanpa_ou
from account_bank_statement_line
"""

rows = q(UNMAPPED)
tot = q(TOTALS)[0]

wb = Workbook()
ws = wb.active
ws.title = "MID Belum Terpetakan"
ws["A1"] = "MID bank yang belum terpetakan ke toko — %s" % DB
ws["A1"].font = Font(bold=True, size=13)
note = ws.cell(
    row=2,
    column=1,
    value=(
        "Satu baris per MID, bukan per mutasi: memetakan satu MID menyelesaikan setiap baris "
        "yang membawanya, lalu dan nanti. Petakan lewat Accounting > Configuration > Bank MID "
        "Mapping, atau langsung dari layar Bank Reconciliation lewat tombol "
        "'Petakan MID ke toko' pada barisnya."
    ),
)
note.font = NOTE
note.alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A2:G2")
ws.row_dimensions[2].height = 32

heads = ["MID", "Jml Baris", "Pertama", "Terakhir", "Nilai", "Jurnal", "Contoh Narasi"]
widths = [18, 11, 13, 13, 18, 22, 72]
for i, (h, w) in enumerate(zip(heads, widths), start=1):
    c = ws.cell(row=4, column=i, value=h)
    c.fill, c.font = HDR, HF
    c.alignment = Alignment(horizontal="center")
    ws.column_dimensions[get_column_letter(i)].width = w

r = 5
total_nilai = 0.0
for row in rows:
    ws.cell(row=r, column=1, value=row["mid"]).fill = WARN
    ws.cell(row=r, column=2, value=int(row["baris"]))
    ws.cell(row=r, column=3, value=row["pertama"])
    ws.cell(row=r, column=4, value=row["terakhir"])
    nilai = float(row["nilai"] or 0)
    total_nilai += nilai
    c = ws.cell(row=r, column=5, value=nilai)
    c.number_format = "#,##0.00"
    ws.cell(row=r, column=6, value=row["jurnal"])
    ws.cell(row=r, column=7, value=row["contoh_narasi"])
    r += 1

ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
ws.cell(row=r, column=2, value=sum(int(x["baris"]) for x in rows)).font = Font(bold=True)
c = ws.cell(row=r, column=5, value=total_nilai)
c.font, c.number_format = Font(bold=True), "#,##0.00"
ws.freeze_panes = "A5"

w2 = wb.create_sheet("Ringkasan")
w2["A1"] = "Posisi pemetaan MID"
w2["A1"].font = Font(bold=True, size=12)
w2.column_dimensions["A"].width = 46
w2.column_dimensions["B"].width = 14
for i, (label, val) in enumerate(
    [
        ("Baris mutasi bank", int(tot["baris"])),
        ("  di antaranya membawa MID", int(tot["ber_mid"])),
        ("  MID terbaca tetapi toko belum ter-resolve", int(tot["mid_tanpa_ou"])),
        ("  tanpa MID dan tanpa toko (perlu baca narasi)", int(tot["tanpa_mid_tanpa_ou"])),
        ("MID unik yang perlu dipetakan", len(rows)),
    ],
    start=3,
):
    w2.cell(row=i, column=1, value=label)
    w2.cell(row=i, column=2, value=val)

wb.save(OUT)
try:
    import grp
    import pwd

    os.chown(OUT, pwd.getpwnam("sftpshare").pw_uid, grp.getgrnam("sftpusers").gr_gid)
except (KeyError, PermissionError, OSError):
    pass

print("Database   : %s" % DB)
print("Baris      : %s  (ber-MID %s)" % (tot["baris"], tot["ber_mid"]))
print("MID tanpa toko : %s baris, %s MID unik, %.2f" % (tot["mid_tanpa_ou"], len(rows), total_nilai))
print("Tersimpan  : %s" % OUT)
