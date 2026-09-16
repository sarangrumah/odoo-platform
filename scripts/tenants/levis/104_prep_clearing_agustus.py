"""Build the August-2026 clearing plan for prd_levis_begbal from EBR's workbook.

Runs on the HOST (plain python3 + openpyxl). Reads only; the database is read
through ``docker exec ... psql`` inside the odoo container, so neither the ORM
nor a database password is needed here. Output:

  * a JSON plan consumed by ``105_clearing_agustus.py`` (odoo shell)
  * a validation workbook for Accounting/FICO

    python3 scripts/tenants/levis/104_prep_clearing_agustus.py \
        --ebr '/path/08. REPORT EBR PERIODE AGUSTUS 2026.xlsx' \
        --db prd_levis_begbal \
        --json /tmp/clearing_agustus.json \
        --xlsx /srv/sftp-share/files/Draft_Clearing_Agustus2026.xlsx

Two blocks only, unlike July:

  A  settlement of August POS receivables (per settlement date, store, bank):
     Dr 1103000002 Bank Suspense + Dr 7104000001 MDR / Cr 1106000101..110
  B  collection in August of the July receivable still open (sheet ``AR JULI``,
     the rows carrying a CASH RECEIVED DATE): same posting, but the credit is
     taken from the *July* POS receivable lines.

There is deliberately no block C: Finance already booked the whole August ATS
sweep by hand (35 GLJV entries, ref "ATS ...", Rp 13.918.328.412,70), so
building it again would double the sweep. There is no block S either: the
August bank statements are complete in Odoo (IBCA 2.038 lines, IBRI 306).

Grain note: the credit side comes from ``COMPILE SALES`` (transaction level,
ties to Odoo per store x trans-date x tender), the bank side from the ``MUTASI``
sheets. MDR is taken per store per bank for the whole month and spread over that
store's settlement dates pro-rata to gross, which keeps every entry balanced and
the monthly MDR per store exact.
"""

import argparse
import collections
import json
import re
import subprocess
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

MONTH_FROM = "2026-08-01"
MONTH_TO = "2026-08-31"
JULY_FROM = "2026-07-01"
JULY_TO = "2026-07-31"

ACC_SUSPENSE = "1103000002"
ACC_MDR = "7104000001"

TENDER_ACC = {
    "CASH": "1106000101",
    "OFFLINE_DOMESTIC_CARD": "1106000102",
    "OFFLINE_VISA": "1106000103",
    "OFFLINE_MASTERCARD": "1106000104",
    "OFFLINE_OTHER_CREDITCARD": "1106000105",
    "OFFLINE_CREDIT_CARD": "1106000106",
    "OFFLINE_JCB": "1106000107",
    "OFFLINE_BRI_CREDIT_CARD": "1106000108",
    "OFFLINE_AMEX": "1106000109",
    "OFFLINE_OVO": "1106000110",
}
# spellings EBR uses that mean the same tender (August adds two more typos)
TENDER_ALIAS = {
    "OFFLINE_DOMESTIC CARD": "OFFLINE_DOMESTIC_CARD",
    "ONLINE_DOMESTIC_CARD": "OFFLINE_DOMESTIC_CARD",
    "OFFLINE_DOMESTIC_CRAD": "OFFLINE_DOMESTIC_CARD",
    "DEBIT": "OFFLINE_DOMESTIC_CARD",
    "OFFLINE_OTHER_CARD": "OFFLINE_OTHER_CREDITCARD",
    "OFFLINE_OTHER_CREDIT_CARD": "OFFLINE_OTHER_CREDITCARD",
    "OFFLINE_BRI_CREDITCARD": "OFFLINE_BRI_CREDIT_CARD",
    "OFFLINE_CREDITCARD": "OFFLINE_CREDIT_CARD",
}

# 0-based row holding the column names -- BCA moved down one row since July
SHEET_HEADER_ROW = {
    "COMPILE SALES": 1,
    "AR JULI": 1,
    "MUTASI BCA": 7,
    "MUTASI BRI": 8,
    "Sheet2": 0,
}


def num(raw):
    s = str(raw or "").replace(",", "").strip()
    if s in ("", "-"):
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def rnd(x):
    return round(x + 0.0, 2)


def sheet_rows(wb, name):
    ws = wb[name]
    rows = [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)]
    hdr = SHEET_HEADER_ROW[name]
    head = [str(c).strip() for c in rows[hdr]]
    return [dict(zip(head, r)) for r in rows[hdr + 1 :] if any(str(c).strip() for c in r)]


def day(raw):
    return str(raw)[:10]


def bank_of_status(status):
    m = re.search(r"\((\w+)\)", str(status or ""))
    return m.group(1) if m else None


def psql(db, sql):
    """Run SQL as the odoo container does -- it already holds the credentials."""
    script = (
        'PGPASSWORD=$(grep -i "^db_password" /etc/odoo/odoo.conf | cut -d= -f2- | tr -d " ") '
        'psql -h postgres -U odoo -d "$1" -At -F "\t" -c "$2"'
    )
    out = subprocess.run(
        ["docker", "exec", "-i", "odoo19-platform-odoo", "sh", "-c", script, "sh", db, sql],
        capture_output=True,
        text=True,
        check=True,
    ).stdout
    return [line.split("\t") for line in out.splitlines() if line]


def store_map(wb):
    """EBR store code -> the analytic-account name Odoo uses."""
    out = {}
    for r in sheet_rows(wb, "Sheet2"):
        code, name = str(r.get("KODE STORE", "")).strip(), str(r.get("NAMA STORE", "")).strip()
        if code and name:
            out[code] = name
    return out


# ------------------------------------------------------------------ bank side
def bank_rows(wb):
    """MDR per (store code, bank), split between settlement and AR collection."""
    mdr_sales = collections.Counter()
    mdr_coll = collections.Counter()
    coll_bank = collections.Counter()
    other = []
    for sheet, bank, dcol in (("MUTASI BCA", "BCA", "Tanggal Transaksi"), ("MUTASI BRI", "BRI", "Tanggal")):
        for r in sheet_rows(wb, sheet):
            code = str(r.get("Store code", "")).strip()
            note = str(r.get("NOTES", "")).strip().upper()
            amount, mdr = num(r.get("AMOUNT PAYMENT")), num(r.get("MDR"))
            if not code or code in ("0", "#N/A"):
                if amount:
                    other.append(
                        {
                            "date": day(r[dcol]),
                            "bank": bank,
                            "flow": str(r.get("CASH IN/ATS", "")).strip(),
                            "desc": str(r.get("Keterangan") or r.get("Uraian Transaksi") or "")[:120],
                            "amount": rnd(amount),
                            "note": note,
                        }
                    )
                continue
            if "COLLECTION" in note:
                mdr_coll[(code, bank)] += mdr
                coll_bank[(code, bank)] += amount
            elif "PELUNASAN" in note:
                other.append(
                    {
                        "date": day(r[dcol]),
                        "bank": bank,
                        "flow": str(r.get("CASH IN/ATS", "")).strip(),
                        "desc": str(r.get("Keterangan") or r.get("Uraian Transaksi") or "")[:120],
                        "amount": rnd(amount),
                        "note": note,
                    }
                )
            else:
                mdr_sales[(code, bank)] += mdr
    return mdr_sales, mdr_coll, coll_bank, other


# ---------------------------------------------------------------- block A / B
def build_block(rows, mdr_month, stores, tag):
    """One row per (settlement date, store, bank) with MDR spread pro-rata.

    `rows` are transaction rows carrying STORE CODE, TRANS DATE, TENDER AMOUNT,
    STATUS and CASH RECEIVED DATE.
    """
    credits = collections.defaultdict(lambda: collections.defaultdict(collections.Counter))
    gross = collections.defaultdict(collections.Counter)
    ebr_split = collections.Counter()
    unsettled = collections.Counter()
    unknown_tender = collections.Counter()
    for r in rows:
        amount = num(r["TENDER AMOUNT"])
        if not amount:
            continue
        code = str(r["STORE CODE"]).strip()
        store = stores.get(code, str(r["STORE NAME"]).strip())
        bank = bank_of_status(r["STATUS"])
        sdate = day(r["CASH RECEIVED DATE"])
        if not bank or not sdate.startswith("2026-08"):
            unsettled[store] += amount
            continue
        tender = str(r["TENDER TYPE"]).strip().upper()
        acc = TENDER_ACC.get(TENDER_ALIAS.get(tender, tender))
        if acc:
            ebr_split[(store, day(r["TRANS DATE"]), acc)] += amount
        else:
            unknown_tender[(store, day(r["TRANS DATE"]), str(r["METODE PEMBAYARAN"]).strip())] += amount
        credits[(code, store, bank)][sdate][day(r["TRANS DATE"])] += amount
        gross[(code, store, bank)][sdate] += amount

    block = []
    for key, per_date in credits.items():
        code, store, bank = key
        total = sum(gross[key].values())
        mdr_total = rnd(mdr_month.get((code, bank), 0.0))
        left = mdr_total
        dates = sorted(per_date)
        for i, sdate in enumerate(dates):
            g = gross[key][sdate]
            mdr = left if i == len(dates) - 1 else (rnd(mdr_total * g / total) if total else 0.0)
            left = rnd(left - mdr)
            block.append(
                {
                    "date": sdate,
                    "store": store,
                    "store_code": code,
                    "bank": bank,
                    "gross": rnd(g),
                    "mdr": rnd(mdr),
                    "cash_in": rnd(g - mdr),
                    "credits": [{"date": d, "amount": rnd(v)} for d, v in sorted(per_date[sdate].items())],
                }
            )
    block.sort(key=lambda r: (r["date"], r["store"], r["bank"]))
    diag = {
        "%s_unsettled" % tag: {k: rnd(v) for k, v in unsettled.items()},
        "%s_unsettled_total" % tag: rnd(sum(unsettled.values())),
        "%s_unknown_tender" % tag: {" | ".join(map(str, k)): rnd(v) for k, v in unknown_tender.items()},
        "%s_ebr_split" % tag: [(s, d, a, rnd(v)) for (s, d, a), v in sorted(ebr_split.items())],
    }
    return block, diag


# ------------------------------------------------------------------- report
def build_workbook(path, plan, diag):
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    money = "#,##0.00"

    def sheet(title, headers, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = bold, head_fill, Alignment(horizontal="center")
        for r in rows:
            ws.append(list(r))
        for i, w in enumerate(widths or [], start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = money
        ws.freeze_panes = "A2"
        return ws

    odoo = plan["odoo"]
    a_gross = sum(r["gross"] for r in plan["block_a"])
    a_mdr = sum(r["mdr"] for r in plan["block_a"])
    b_gross = sum(r["gross"] for r in plan["block_b"])
    b_mdr = sum(r["mdr"] for r in plan["block_b"])

    ws = wb.active
    ws.title = "RINGKASAN"
    rows = [
        ("Clearing Agustus 2026 - prd_levis_begbal", "", ""),
        ("Sumber", "08. REPORT EBR PERIODE AGUSTUS 2026.xlsx", ""),
        ("", "", ""),
        ("BLOK", "URAIAN", "JUMLAH"),
        ("A", "Settlement Agustus - bruto (Cr POS Receivable 1106000101..110)", rnd(a_gross)),
        ("A", "   di antaranya MDR (Dr 7104000001)", rnd(a_mdr)),
        ("A", "   cash-in bersih (Dr 1103000002)", rnd(a_gross - a_mdr)),
        ("B", "Collection AR Juli - bruto (Cr POS Receivable Juli)", rnd(b_gross)),
        ("B", "   di antaranya MDR", rnd(b_mdr)),
        ("B", "   cash-in bersih", rnd(b_gross - b_mdr)),
        ("A+B", "Total MDR Agustus", rnd(a_mdr + b_mdr)),
        ("", "", ""),
        ("Odoo", "Penjualan Agustus (POS receivable dibuka RIREC)", odoo["sales_aug"]),
        ("Odoo", "POS Receivable Agustus masih open", odoo["posrec_aug"]),
        ("Odoo", "   sisa setelah blok A", rnd(odoo["posrec_aug"] - a_gross)),
        ("Odoo", "POS Receivable Juli masih open", odoo["posrec_jul"]),
        ("Odoo", "   sisa setelah blok B", rnd(odoo["posrec_jul"] - b_gross)),
        ("Odoo", "Saldo Bank Suspense 1103000002 sebelum clearing", odoo["suspense"]),
        ("Odoo", "Saldo MDR 7104000001 sebelum clearing", odoo["mdr"]),
        ("", "", ""),
        ("SIMULASI SETELAH BLOK A+B DIPOSTING (dry-run 105_clearing_agustus.py)", "", ""),
        ("Odoo", "Bank Suspense 1103000002 per 31-Agu", odoo["sim_suspense"]),
        ("Odoo", "MDR 7104000001 kumulatif per 31-Agu", odoo["sim_mdr"]),
        ("Odoo", "POS Receivable Agustus tersisa", odoo["sim_posrec_aug"]),
        ("Odoo", "POS Receivable Juli tersisa", odoo["sim_posrec_jul"]),
        ("", "", ""),
        ("Catatan", "Tender belum tersettle per 31-Agu (tetap open)", diag["a_unsettled_total"]),
        (
            "Catatan",
            "Sisa AR Juli yang TIDAK ditagih Agustus (KOL Grand Indonesia + voucher Juni)",
            diag["b_unsettled_total"],
        ),
        ("Catatan", "Sweep ATS Agustus SUDAH dijurnal manual Finance - tidak dibuat ulang di sini", odoo["ats_manual"]),
        ("", "", ""),
        ("PERLU KEPUTUSAN ACCOUNTING", "", ""),
        ("1", "Mandiri & BNI tidak masuk clearing Agustus (dijurnal manual klien)", ""),
        ("2", "Selisih sisa AR Juli antara buku Odoo dan sheet AR JULI EBR", odoo["ar_juli_gap"]),
        ("3", "Baris bank tanpa toko (PELUNASAN QR, bunga, pajak, biaya admin) - di luar clearing", ""),
    ]
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=13)
    for row in ws.iter_rows(min_row=5):
        if isinstance(row[2].value, float):
            row[2].number_format = money
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 20

    for tag, title, block in (
        ("a", "A - SETTLEMENT AGUSTUS", plan["block_a"]),
        ("b", "B - COLLECTION AR JULI", plan["block_b"]),
    ):
        sheet(
            title,
            ["Tgl settle", "Toko", "Bank", "Bruto", "MDR", "Cash in", "Jml baris kredit"],
            [(r["date"], r["store"], r["bank"], r["gross"], r["mdr"], r["cash_in"], len(r["credits"])) for r in block],
            [12, 38, 8, 18, 14, 18, 16],
        )
        sheet(
            title.split(" - ")[0] + " - RINCIAN KREDIT",
            ["Tgl settle", "Toko", "Bank", "Tgl transaksi", "Jumlah"],
            [(r["date"], r["store"], r["bank"], c["date"], c["amount"]) for r in block for c in r["credits"]],
            [12, 38, 8, 14, 18],
        )
    sheet(
        "REKON PENJUALAN HARIAN",
        ["Tgl transaksi", "EBR", "Odoo", "Selisih"],
        plan["recon_daily"],
        [14, 20, 20, 16],
    )
    sheet(
        "REKON PIUTANG JULI",
        ["Toko", "Open di Odoo", "Ditagih Agustus", "Sisa menurut EBR", "Selisih"],
        plan["recon_july"],
        [38, 18, 18, 18, 16],
    )
    sheet(
        "MDR PER TOKO",
        ["Toko", "Bank", "MDR settlement", "MDR collection"],
        plan["mdr_table"],
        [38, 8, 18, 18],
    )
    sheet(
        "BARIS BANK TANPA TOKO",
        ["Tanggal", "Bank", "Arus", "Catatan", "Keterangan", "Jumlah"],
        [(r["date"], r["bank"], r["flow"], r["note"], r["desc"], r["amount"]) for r in plan["bank_other"]],
        [12, 8, 14, 22, 70, 18],
    )
    sheet(
        "BELUM TERSETTLE 31-AGU",
        ["Toko", "Jumlah"],
        sorted(diag["a_unsettled"].items(), key=lambda kv: -kv[1]),
        [38, 18],
    )
    if diag["a_unknown_tender"]:
        sheet(
            "TENDER TIDAK DIKENAL",
            ["Toko | Tanggal | Metode", "Jumlah"],
            sorted(diag["a_unknown_tender"].items(), key=lambda kv: -kv[1]),
            [64, 18],
        )
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ebr", required=True)
    ap.add_argument("--db", default="prd_levis_begbal")
    ap.add_argument("--json", required=True)
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--sim-suspense", type=float, default=0.0)
    ap.add_argument("--sim-mdr", type=float, default=0.0)
    ap.add_argument("--sim-posrec-aug", type=float, default=0.0)
    ap.add_argument("--sim-posrec-jul", type=float, default=0.0)
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.ebr, read_only=True, data_only=True)
    stores = store_map(wb)
    mdr_sales, mdr_coll, coll_bank, bank_other = bank_rows(wb)

    block_a, diag_a = build_block(sheet_rows(wb, "COMPILE SALES"), mdr_sales, stores, "a")
    block_b, diag_b = build_block(sheet_rows(wb, "AR JULI"), mdr_coll, stores, "b")
    diag = dict(diag_a, **diag_b)

    # the collection rows must tie to the bank, per store and bank
    for key, amount in coll_bank.items():
        booked = sum(r["gross"] for r in block_b if (r["store_code"], r["bank"]) == key)
        if abs(booked - amount) > 0.5:
            print("WARN collection %s: sheet AR JULI %.2f vs mutasi bank %.2f" % (key, booked, amount))

    known = {r[0] for r in psql(args.db, "select name->>'en_US' from account_analytic_account")}
    unknown = sorted({r["store"] for r in block_a + block_b} - known)
    if unknown:
        sys.exit("store names with no analytic account: " + ", ".join(unknown))

    # ---------------------------------------------------------------- Odoo side
    def one(sql):
        return float(psql(args.db, sql)[0][0] or 0)

    posrec = """
        select coalesce(sum(aml.amount_residual), 0)
        from account_move_line aml join account_account aa on aa.id = aml.account_id
        where aml.parent_state = 'posted' and not aml.reconciled
          and aa.code_store->>'1' between '1106000101' and '1106000110'
          and aml.date between '%s' and '%s'"""
    odoo = {
        "posrec_aug": rnd(one(posrec % (MONTH_FROM, MONTH_TO))),
        "posrec_jul": rnd(one(posrec % (JULY_FROM, JULY_TO))),
        "sales_aug": rnd(
            one(
                """select coalesce(sum(aml.debit), 0)
                   from account_move_line aml join account_account aa on aa.id = aml.account_id
                   where aml.parent_state = 'posted'
                     and aa.code_store->>'1' between '1106000101' and '1106000110'
                     and aml.date between '%s' and '%s'"""
                % (MONTH_FROM, MONTH_TO)
            )
        ),
        "ats_manual": rnd(
            one(
                """select coalesce(sum(aml.debit), 0)
                   from account_move_line aml
                   join account_move am on am.id = aml.move_id
                   join account_account aa on aa.id = aml.account_id
                   where am.state = 'posted' and am.ref like 'ATS%%'
                     and aa.code_store->>'1' = '1103019320'
                     and am.date between '%s' and '%s'"""
                % (MONTH_FROM, MONTH_TO)
            )
        ),
    }
    bal = {
        r[0]: rnd(float(r[1]))
        for r in psql(
            args.db,
            """select aa.code_store->>'1', sum(aml.debit - aml.credit)
               from account_move_line aml join account_account aa on aa.id = aml.account_id
               where aml.parent_state = 'posted' and aml.date <= '%s'
                 and aa.code_store->>'1' in ('%s', '%s') group by 1"""
            % (MONTH_TO, ACC_SUSPENSE, ACC_MDR),
        )
    }
    odoo["sim_suspense"] = rnd(args.sim_suspense)
    odoo["sim_mdr"] = rnd(args.sim_mdr)
    odoo["sim_posrec_aug"] = rnd(args.sim_posrec_aug)
    odoo["sim_posrec_jul"] = rnd(args.sim_posrec_jul)
    odoo["suspense"] = bal.get(ACC_SUSPENSE, 0.0)
    odoo["mdr"] = bal.get(ACC_MDR, 0.0)

    # daily sales EBR vs Odoo
    odoo_daily = {
        r[0]: rnd(float(r[1]))
        for r in psql(
            args.db,
            """select aml.date::text, sum(aml.debit)
               from account_move_line aml join account_account aa on aa.id = aml.account_id
               where aml.parent_state = 'posted'
                 and aa.code_store->>'1' between '1106000101' and '1106000110'
                 and aml.date between '%s' and '%s' group by 1"""
            % (MONTH_FROM, MONTH_TO),
        )
    }
    ebr_daily = collections.Counter()
    for r in sheet_rows(wb, "COMPILE SALES"):
        ebr_daily[day(r["TRANS DATE"])] += num(r["TENDER AMOUNT"])
    recon_daily = [
        (d, rnd(ebr_daily.get(d, 0.0)), odoo_daily.get(d, 0.0), rnd(ebr_daily.get(d, 0.0) - odoo_daily.get(d, 0.0)))
        for d in sorted(set(ebr_daily) | set(odoo_daily))
    ]

    # July receivable: Odoo open vs what August collects
    odoo_july = collections.Counter()
    for r in psql(
        args.db,
        """select an.name->>'en_US', sum(aml.amount_residual)
           from account_move_line aml
           join account_account aa on aa.id = aml.account_id
           join account_analytic_account an
             on an.id = (select (jsonb_object_keys(aml.analytic_distribution))::int limit 1)
           where aml.parent_state = 'posted' and not aml.reconciled
             and aa.code_store->>'1' between '1106000101' and '1106000110'
             and aml.date between '%s' and '%s' group by 1"""
        % (JULY_FROM, JULY_TO),
    ):
        odoo_july[r[0]] += rnd(float(r[1]))
    coll_by_store = collections.Counter()
    for r in block_b:
        coll_by_store[r["store"]] += r["gross"]
    ebr_left = collections.Counter()
    for store, amount in diag["b_unsettled"].items():
        ebr_left[store] += amount
    recon_july = []
    for store in sorted(set(odoo_july) | set(coll_by_store) | set(ebr_left)):
        o, c, l = odoo_july.get(store, 0.0), coll_by_store.get(store, 0.0), ebr_left.get(store, 0.0)
        recon_july.append((store, rnd(o), rnd(c), rnd(l), rnd(o - c - l)))
    odoo["ar_juli_gap"] = rnd(sum(r[4] for r in recon_july))

    inv = {v: k for k, v in stores.items()}
    mdr_table = sorted(
        {
            (stores.get(code, code), bank, rnd(mdr_sales.get((code, bank), 0.0)), rnd(mdr_coll.get((code, bank), 0.0)))
            for code, bank in set(mdr_sales) | set(mdr_coll)
        }
    )

    plan = {
        "month": MONTH_FROM[:7],
        "block_a": block_a,
        "block_b": block_b,
        "bank_other": bank_other,
        "recon_daily": recon_daily,
        "recon_july": recon_july,
        "mdr_table": mdr_table,
        "odoo": odoo,
    }
    with open(args.json, "w") as fh:
        json.dump(plan, fh)
    build_workbook(args.xlsx, plan, diag)

    a_gross = sum(r["gross"] for r in block_a)
    b_gross = sum(r["gross"] for r in block_b)
    print(f"block A  {len(block_a):5d} rows  gross {a_gross:>20,.2f}  mdr {sum(r['mdr'] for r in block_a):>16,.2f}")
    print(f"block B  {len(block_b):5d} rows  gross {b_gross:>20,.2f}  mdr {sum(r['mdr'] for r in block_b):>16,.2f}")
    print(f"unsettled 31-Agu               {diag['a_unsettled_total']:>20,.2f}")
    print(f"AR Juli tidak ditagih          {diag['b_unsettled_total']:>20,.2f}")
    print(
        f"POS receivable Agu (Odoo)      {odoo['posrec_aug']:>20,.2f}  -> after A {odoo['posrec_aug'] - a_gross:>18,.2f}"
    )
    print(
        f"POS receivable Jul (Odoo)      {odoo['posrec_jul']:>20,.2f}  -> after B {odoo['posrec_jul'] - b_gross:>18,.2f}"
    )
    print(f"selisih AR Juli Odoo vs EBR    {odoo['ar_juli_gap']:>20,.2f}")
    print(f"json {args.json}\nxlsx {args.xlsx}")


if __name__ == "__main__":
    main()
