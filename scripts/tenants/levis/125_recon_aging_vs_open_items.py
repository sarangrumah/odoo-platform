# -*- coding: utf-8 -*-
"""Rekonsiliasi Aged Receivable vs GL Open Items, per tanggal posisi.

Sheet baris #73. SELECT-only -- tidak pernah menulis apa pun.

  docker exec -i -e OUT=/mnt/data_levis/Rekonsiliasi_Aging_vs_OpenItems.xlsx \\
    -e AS_OF=2026-08-31 odoo19-platform-odoo \\
    odoo shell -d prd_levis_begbal --no-http --max-cron-threads=0 \\
    --shell-interface=python \\
    < scripts/tenants/levis/125_recon_aging_vs_open_items.py

--------------------------------------------------------------------------
Pertanyaan klien, dan jawabannya
--------------------------------------------------------------------------
"baris transaksi outstanding aging report vs open items report berbeda jauh."

Benar, dan itu **bukan selisih uang**. Diukur di ``prd_levis_begbal``
per 31-Agu-2026:

    Aged Receivable   1.331 baris dokumen   Rp 3.088.025.495
    GL Open Items       114 baris           Rp 3.088.025.495

Totalnya **ikat sampai rupiah**. Yang berbeda adalah pertanyaan yang dijawab:

* **Aged Receivable** = daftar kerja penagihan. Satu baris per **dokumen** yang
  masih terbuka pada tanggal posisi, dikelompokkan per partner.
* **GL Open Items** = daftar kerja clearing. Satu baris per **sisa yang belum
  saling menutup** setelah FIFO netting per (akun, partner) -- jadi debit dan
  kredit yang sudah saling meniadakan hilang dari daftar, dan 1.331 baris
  mengerut jadi 114.

Ada satu sebab tambahan yang membuat selisih barisnya terasa ekstrem di tenant
ini: **seluruh 1.331 baris receivable terbuka TIDAK punya partner** (POS
settlement dan jurnal bank). Aged Receivable karena itu menampilkan satu grup
tunggal "No Partner", sementara netting GL Open Items jadi seagresif mungkin
karena semua baris berbagi kunci partner yang sama (kosong).

Laporan ini menuliskan kedua sisi berdampingan supaya Accounting bisa melihat
sendiri bahwa uangnya sama.
"""

import os
from collections import defaultdict
from datetime import date as date_cls

OUT = os.environ.get("OUT", "/mnt/data_levis/Rekonsiliasi_Aging_vs_OpenItems.xlsx")
AS_OF = os.environ.get("AS_OF", "2026-08-31")
KIND = os.environ.get("KIND", "receivable")  # receivable | payable

as_of = date_cls.fromisoformat(AS_OF)
companies = env["res.company"].search([]).ids
company = env["res.company"].browse(companies[0])
currency = company.currency_id

aged_model = "custom.report.aged.%s" % KIND
account_type = "asset_receivable" if KIND == "receivable" else "liability_payable"

print("=" * 78)
print("Rekonsiliasi Aged %s vs GL Open Items   per %s" % (KIND.title(), AS_OF))
print("=" * 78)

aged = env[aged_model]
aged_opts = {
    "date_from": date_cls(1970, 1, 1),
    "date_to": as_of,
    "company_ids": companies,
    "partner_ids": [],
    "posted_only": True,
}
summary = aged._build_summary_lines(aged_opts)
detail = aged._build_detail_lines({**aged_opts, "layout": "detail"})
groups = detail.get("partners", [])
doc_rows = [row for group in groups for row in group.get("rows", [])]
aged_total = summary["grand_total"]["total"]

print("\nAged %s" % KIND.title())
print("  grup partner        : %s" % len(summary["rows"]))
print("  baris dokumen       : %s" % len(doc_rows))
print("  total outstanding   : %s" % currency.round(aged_total))
no_partner = [r for r in summary["rows"] if (r.get("partner_name") or "") == "No Partner"]
if no_partner:
    print("  di antaranya 'No Partner': %s" % currency.round(no_partner[0]["total"]))

gl = env["custom.report.gl.open.items"]
gl_lines = gl._build_lines(
    {
        "date_from": None,
        "date_to": as_of,
        "company_ids": companies,
        "partner_ids": [],
        "account_ids": [],
        "account_types": [account_type],
    }
)
gl_rows = [line for line in gl_lines if not line.get("type")]
gl_total = sum(line.get("outstanding") or 0.0 for line in gl_rows)

print("\nGL Open Items")
print("  baris (sesudah netting): %s" % len(gl_rows))
print("  total outstanding      : %s" % currency.round(gl_total))

gap = aged_total - gl_total
print("\nSELISIH UANG: %s" % currency.round(gap))
if currency.is_zero(gap):
    print("  -> NOL. Kedua laporan memuat uang yang sama; yang berbeda hanya bentuk barisnya.")
else:
    print("  -> TIDAK NOL. Ini perlu ditelusuri, bukan sekadar perbedaan bentuk.")
print("SELISIH BARIS: %s dokumen vs %s baris netting" % (len(doc_rows), len(gl_rows)))

per_account_gl = defaultdict(lambda: [0, 0.0])
for row in gl_rows:
    bucket = per_account_gl[row.get("account") or "?"]
    bucket[0] += 1
    bucket[1] += row.get("outstanding") or 0.0

per_account_aged = defaultdict(lambda: [0, 0.0])
for row in doc_rows:
    bucket = per_account_aged[row.get("account") or row.get("account_name") or "?"]
    bucket[0] += 1
    bucket[1] += row.get("outstanding") or row.get("total") or 0.0

print("\nPer akun (GL Open Items):")
for account, (count, amount) in sorted(per_account_gl.items()):
    print("  %-46s %4s baris  %s" % (account[:46], count, currency.round(amount)))

try:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    ws.append(["Tanggal posisi", AS_OF])
    ws.append([])
    ws.append(["Laporan", "Jumlah baris", "Total outstanding", "Satu baris artinya"])
    ws.append(
        [
            "Aged %s (detail)" % KIND.title(),
            len(doc_rows),
            aged_total,
            "satu dokumen yang masih terbuka pada tanggal posisi",
        ]
    )
    ws.append(["GL Open Items", len(gl_rows), gl_total, "satu sisa yang belum saling menutup sesudah FIFO netting"])
    ws.append([])
    ws.append(["Selisih uang", "", gap, "nol = kedua laporan memuat uang yang sama"])
    ws.append([])
    ws.append(["Catatan"])
    ws.append(["Seluruh baris terbuka di tenant ini tidak punya partner (POS settlement, jurnal bank),"])
    ws.append(["sehingga Aged menampilkan satu grup 'No Partner' dan netting GL Open Items menjadi"])
    ws.append(["seagresif mungkin. Itu sebabnya selisih barisnya terasa ekstrem."])

    w2 = wb.create_sheet("GL Open Items")
    w2.append(["Akun", "Baris", "Outstanding"])
    for account, (count, amount) in sorted(per_account_gl.items()):
        w2.append([account, count, amount])

    w3 = wb.create_sheet("Aged detail")
    if doc_rows:
        keys = [k for k in doc_rows[0].keys() if not k.startswith("_")]
        w3.append(keys)
        for row in doc_rows:
            w3.append([row.get(k) for k in keys])
    else:
        w3.append(["(tidak ada baris terbuka)"])

    wb.save(OUT)
    print("\nLaporan: %s" % OUT)
except Exception as exc:  # noqa: BLE001
    print("\n(laporan xlsx gagal: %s)" % str(exc)[:80])

env.cr.rollback()
