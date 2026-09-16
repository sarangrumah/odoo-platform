# Konfirmasi PPN Masukan PO EBR -- "Included" vs "Excluded", sample PO/T/EBR/2026/08/00075.
#
# Latar: informasi awal ke tim adalah harga di PO sudah termasuk PPN (included).
# Ternyata harga itu DPP (excluded) -- PPN 12% (efektif 11%) harus ditambahkan di atasnya.
# Di prd_levis_begbal semua 199 PO punya amount_tax = 0 karena kolom Taxes dikosongkan
# saat entry, sedangkan default pajak produk (tax id 21) mewarisi
# res_company.account_price_include = 'tax_included' -- kalau default itu dipakai apa
# adanya, Rp 231.882.120 justru diurai jadi DPP 208.902.810,81. Standar rumah adalah
# tax id 38 "PPN 12% (Excluded)" -- itu yang dipakai di seluruh tagihan yang sudah ada.
#
# Script ini SELECT-ONLY. Tidak ada satu pun UPDATE/INSERT ke Odoo; satu-satunya tulisan
# adalah file Excel di /srv/sftp-share/files (bisa diunduh lewat File Browser /files).
#
#   python3 scripts/tenants/levis/83_report_po_ppn_excluded.py
#
# Env:  PO_NAME  -> nomor PO sample (default PO/T/EBR/2026/08/00075)
#       DB       -> database (default prd_levis_begbal)
#       OUT      -> path file xlsx

import csv
import io
import os
import subprocess
import sys
from decimal import ROUND_HALF_UP, Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PG = "odoo19-platform-postgres"
DB = os.environ.get("DB", "prd_levis_begbal")
PO_NAME = os.environ.get("PO_NAME", "PO/T/EBR/2026/08/00075")
OUT = os.environ.get("OUT", "/srv/sftp-share/files/Konfirmasi_PPN_PO-T-EBR-2026-08-00075.xlsx")

# tax id 38 "PPN 12% (Excluded)": amount 11,0 / price_include_override tax_excluded /
# dpp_method regular. PPN 12% atas DPP nilai lain 11/12 => tarif efektif 11% dari harga.
RATE = Decimal("0.11")
CENT = Decimal("0.01")

MONEY = "#,##0.00"
QTY = "#,##0"
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF")
SUB_FONT = Font(bold=True)
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")


def q(sql):
    """Run a read-only query and return a list of dicts."""
    out = subprocess.run(
        [
            "docker",
            "exec",
            PG,
            "sh",
            "-c",
            f'PGPASSWORD="$POSTGRES_PASSWORD" psql -U "$POSTGRES_USER" -d {DB} --csv -v ON_ERROR_STOP=1 -c "{sql}"',
        ],
        capture_output=True,
        text=True,
    )
    if out.returncode:
        sys.exit(f"query failed:\n{out.stderr}")
    return list(csv.DictReader(io.StringIO(out.stdout)))


def money(v):
    return Decimal(v).quantize(CENT, rounding=ROUND_HALF_UP)


def f(v):
    """Excel wants a float, but every number above is computed in Decimal first."""
    return float(v)


def head(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


# --------------------------------------------------------------------------- data
po = q(
    "SELECT po.id, po.name, po.state, po.date_order::date AS date_order, "
    "po.partner_ref, po.amount_untaxed, po.amount_tax, po.amount_total, "
    "rp.name AS vendor "
    "FROM purchase_order po JOIN res_partner rp ON rp.id = po.partner_id "
    f"WHERE po.name = '{PO_NAME}'"
)
if not po:
    sys.exit(f"PO {PO_NAME} tidak ditemukan di {DB}")
po = po[0]

gr = q(
    "SELECT DISTINCT sp.name, sp.state FROM stock_picking sp "
    "JOIN stock_move sm ON sm.picking_id = sp.id "
    "JOIN purchase_order_line l ON l.id = sm.purchase_line_id "
    f"WHERE l.order_id = {po['id']} ORDER BY 1"
)
gr_txt = ", ".join(f"{p['name']} ({p['state']})" for p in gr) or "-"

lines = q(
    "SELECT pp.default_code, pt.name->>'en_US' AS product, l.name AS descr, "
    "l.product_qty, l.price_unit, l.price_subtotal "
    "FROM purchase_order_line l "
    "JOIN product_product pp ON pp.id = l.product_id "
    "JOIN product_template pt ON pt.id = pp.product_tmpl_id "
    f"WHERE l.order_id = {po['id']} ORDER BY l.id"
)

all_po = q(
    "SELECT po.name, po.date_order::date AS date_order, rp.name AS vendor, "
    "po.amount_untaxed, po.amount_tax "
    "FROM purchase_order po JOIN res_partner rp ON rp.id = po.partner_id "
    "WHERE po.state = 'purchase' AND po.amount_tax = 0 ORDER BY po.name"
)

dpp = money(po["amount_untaxed"])
ppn_excl = money(dpp * RATE)
total_excl = dpp + ppn_excl
# skenario "included": harga dianggap gross, DPP diurai dari dalamnya
dpp_incl = money(dpp / (Decimal(1) + RATE))
ppn_incl = dpp - dpp_incl
qty_total = sum(Decimal(r["product_qty"]) for r in lines)

# --------------------------------------------------------------------- sheet 1
wb = Workbook()
ws = wb.active
ws.title = "Ringkasan"
ws.column_dimensions["A"].width = 38
for col in "BCD":
    ws.column_dimensions[col].width = 22

ws["A1"] = "Konfirmasi Perhitungan PPN Masukan — Purchase Order"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = (
    "Koreksi asumsi: harga PO semula diinformasikan SUDAH termasuk PPN (included); "
    "ternyata harga tersebut adalah DPP (excluded)."
)
ws["A2"].font = Font(italic=True)

r = 4
for label, value in [
    ("Database", DB),
    ("Nomor PO", po["name"]),
    ("Vendor", po["vendor"]),
    ("Tanggal PO", po["date_order"]),
    ("No. Faktur / Referensi Vendor", po["partner_ref"] or "-"),
    ("Status PO", po["state"]),
    ("Penerimaan Barang (GR)", gr_txt),
    ("Status Tagihan", "Belum ditagih — belum ada vendor bill"),
    ("Jumlah Baris / Qty", f"{len(lines)} baris / {qty_total:,.0f} pcs"),
    ("Pajak terpasang di baris PO", "TIDAK ADA (amount_tax = 0)"),
]:
    ws.cell(row=r, column=1, value=label).font = SUB_FONT
    ws.cell(row=r, column=2, value=value)
    r += 1

r += 1
ws.cell(row=r, column=1, value="PERBANDINGAN PERHITUNGAN").font = Font(bold=True, size=12)
r += 1
head(ws, r, ["Skenario", "DPP (Rp)", "PPN Masukan (Rp)", "Total ke Vendor (Rp)"])
r += 1

rows = [
    ("A. Kondisi sekarang di Odoo (tanpa pajak)", dpp, Decimal(0), dpp, WARN_FILL),
    ("B. Asumsi awal — harga INCLUDED PPN", dpp_incl, ppn_incl, dpp, WARN_FILL),
    ("C. Yang benar — harga EXCLUDED PPN", dpp, ppn_excl, total_excl, OK_FILL),
]
for label, a, b, c, fill in rows:
    ws.cell(row=r, column=1, value=label).fill = fill
    for col, val in ((2, a), (3, b), (4, c)):
        cell = ws.cell(row=r, column=col, value=f(val))
        cell.number_format = MONEY
        cell.fill = fill
    r += 1

r += 1
ws.cell(row=r, column=1, value="Selisih B (included) vs C (excluded)").font = SUB_FONT
for col, val in ((2, dpp - dpp_incl), (3, ppn_excl - ppn_incl), (4, total_excl - dpp)):
    cell = ws.cell(row=r, column=col, value=f(val))
    cell.number_format = MONEY
    cell.font = SUB_FONT
r += 2

for note in [
    "CATATAN",
    "1. Tarif: PPN 12% dengan DPP Nilai Lain 11/12 → tarif efektif 11% dari harga "
    '(pajak "PPN 12% (Excluded)" di Odoo).',
    "2. Harga satuan di PO TIDAK berubah. Yang bertambah hanya baris PPN di atas harga.",
    f"3. Karena harga satuan tidak berubah, nilai persediaan dan GR/IR (Rp {dpp:,.2f}) juga TIDAK berubah.",
    "4. PO ini belum pernah ditagih, jadi koreksi tidak memerlukan pembatalan atau reversal jurnal.",
    "5. Pembulatan mengikuti setelan perusahaan round_globally, mata uang IDR 2 desimal.",
    "6. Dokumen ini bersifat KONFIRMASI. Belum ada data apa pun yang diubah di Odoo.",
]:
    ws.cell(row=r, column=1, value=note)
    if note == "CATATAN":
        ws.cell(row=r, column=1).font = Font(bold=True, size=12)
    r += 1

# --------------------------------------------------------------------- sheet 2
ws2 = wb.create_sheet("Detail Baris")
head(
    ws2,
    1,
    [
        "No",
        "Kode Produk",
        "Nama Produk",
        "Qty",
        "Harga Satuan / DPP (Rp)",
        "Subtotal DPP (Rp)",
        "PPN 11% (Rp)",
        "Total termasuk PPN (Rp)",
        "Jika INCLUDED — DPP (Rp)",
        "Jika INCLUDED — PPN (Rp)",
    ],
    widths=[5, 20, 42, 8, 20, 18, 16, 20, 20, 18],
)
ws2.freeze_panes = "A2"

r = 2
sum_sub = sum_ppn = sum_tot = sum_isub = sum_ippn = Decimal(0)
for i, ln in enumerate(lines, start=1):
    sub = money(ln["price_subtotal"])
    ppn = money(sub * RATE)
    tot = sub + ppn
    isub = money(sub / (Decimal(1) + RATE))
    ippn = sub - isub
    sum_sub += sub
    sum_ppn += ppn
    sum_tot += tot
    sum_isub += isub
    sum_ippn += ippn

    ws2.cell(row=r, column=1, value=i)
    ws2.cell(row=r, column=2, value=ln["default_code"])
    ws2.cell(row=r, column=3, value=ln["product"])
    c = ws2.cell(row=r, column=4, value=f(Decimal(ln["product_qty"])))
    c.number_format = QTY
    for col, val in (
        (5, money(ln["price_unit"])),
        (6, sub),
        (7, ppn),
        (8, tot),
        (9, isub),
        (10, ippn),
    ):
        cell = ws2.cell(row=r, column=col, value=f(val))
        cell.number_format = MONEY
    r += 1

ws2.cell(row=r, column=3, value="TOTAL").font = SUB_FONT
c = ws2.cell(row=r, column=4, value=f(qty_total))
c.number_format, c.font = QTY, SUB_FONT
for col, val in (
    (6, sum_sub),
    (7, sum_ppn),
    (8, sum_tot),
    (9, sum_isub),
    (10, sum_ippn),
):
    cell = ws2.cell(row=r, column=col, value=f(val))
    cell.number_format, cell.font = MONEY, SUB_FONT

# Pembulatan per-baris bisa meleset beberapa sen dari pembulatan tingkat dokumen
# (round_globally). Angka yang dipakai untuk konfirmasi adalah yang di Ringkasan.
if sum_ppn != ppn_excl:
    ws2.cell(
        row=r + 2,
        column=3,
        value=(
            f"Catatan: penjumlahan PPN per baris Rp {sum_ppn:,.2f} vs pembulatan tingkat "
            f"dokumen Rp {ppn_excl:,.2f} (selisih Rp {sum_ppn - ppn_excl:,.2f}). "
            "Odoo memakai round_globally, jadi yang berlaku adalah angka dokumen."
        ),
    ).font = Font(italic=True)

# --------------------------------------------------------------------- sheet 3
ws3 = wb.create_sheet("Dampak Seluruh PO EBR")
ws3.cell(
    row=1,
    column=1,
    value=(
        f"Seluruh PO berstatus Purchase Order di {DB} yang saat ini tanpa PPN "
        "(informasi saja — belum ada perubahan data)"
    ),
).font = Font(bold=True, size=12)
head(
    ws3,
    3,
    [
        "No",
        "Nomor PO",
        "Tanggal",
        "Vendor",
        "DPP Sekarang (Rp)",
        "PPN 11% Seharusnya (Rp)",
        "Total Setelah Koreksi (Rp)",
    ],
    widths=[5, 26, 13, 40, 20, 22, 24],
)
ws3.freeze_panes = "A4"

r = 4
t_dpp = t_ppn = Decimal(0)
for i, p in enumerate(all_po, start=1):
    d = money(p["amount_untaxed"])
    v = money(d * RATE)
    t_dpp += d
    t_ppn += v
    ws3.cell(row=r, column=1, value=i)
    ws3.cell(row=r, column=2, value=p["name"])
    ws3.cell(row=r, column=3, value=p["date_order"])
    ws3.cell(row=r, column=4, value=p["vendor"])
    for col, val in ((5, d), (6, v), (7, d + v)):
        cell = ws3.cell(row=r, column=col, value=f(val))
        cell.number_format = MONEY
    r += 1

ws3.cell(row=r, column=4, value=f"TOTAL ({len(all_po)} PO)").font = SUB_FONT
for col, val in ((5, t_dpp), (6, t_ppn), (7, t_dpp + t_ppn)):
    cell = ws3.cell(row=r, column=col, value=f(val))
    cell.number_format, cell.font = MONEY, SUB_FONT

wb.save(OUT)

print(f"PO            : {po['name']}  ({po['vendor']})")
print(f"Baris / qty   : {len(lines)} / {qty_total:,.0f}")
print(f"A. sekarang   : DPP {dpp:>18,.2f}  PPN {Decimal(0):>16,.2f}  Total {dpp:>18,.2f}")
print(f"B. included   : DPP {dpp_incl:>18,.2f}  PPN {ppn_incl:>16,.2f}  Total {dpp:>18,.2f}")
print(f"C. excluded   : DPP {dpp:>18,.2f}  PPN {ppn_excl:>16,.2f}  Total {total_excl:>18,.2f}")
print(f"Seluruh PO    : {len(all_po)} PO, DPP {t_dpp:,.2f}, PPN seharusnya {t_ppn:,.2f}")
print(f"Tersimpan     : {OUT}")
