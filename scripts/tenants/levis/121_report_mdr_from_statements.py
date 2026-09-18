# -*- coding: utf-8 -*-
"""Laporan MDR dari mutasi bank yang sudah terupload.

Sheet baris #50. SELECT-only -- tidak pernah menulis apa pun.

  docker exec -i -e OUT=/mnt/data_levis/Laporan_MDR.xlsx \\
    -e DATE_FROM=2026-07-01 -e DATE_TO=2026-09-30 odoo19-platform-odoo \\
    odoo shell -d prd_levis_begbal --no-http --max-cron-threads=0 \\
    --shell-interface=python < scripts/tenants/levis/121_report_mdr_from_statements.py

--------------------------------------------------------------------------
Dari mana angkanya
--------------------------------------------------------------------------
Tidak ada tabel MDR di Odoo ini. ``levis.mdr.bin`` kosong dan tidak pernah
diisi. MDR yang dipakai sistem **dibaca dari narasi mutasi bank** oleh
``levis.bank.narrative`` dan disimpan di ``account.bank.statement.line.levis_mdr``,
lalu dibukukan ``levis.pos.clearing`` ke akun **7104000001**.

Jadi laporan ini punya dua sisi yang harus dibandingkan:

* **MDR narasi** -- apa yang tertulis di rekening koran (sumber kebenaran bank);
* **MDR dibukukan** -- mutasi GL 7104000001 pada periode yang sama.

Selisih keduanya bukan kesalahan laporan: ia berarti ada baris bank yang belum
masuk clearing, atau clearing yang memakai angka lain.

--------------------------------------------------------------------------
Jebakan yang sudah menggigit sebelumnya
--------------------------------------------------------------------------
* **Impor statement ganda.** IBCA Agustus pernah diimpor kumulatif sehingga
  baris yang sama masuk dua kali. Sheet "Dugaan Duplikat" menandai baris dengan
  (jurnal, tanggal, jumlah, narasi) yang identik -- periksa sebelum memakai
  totalnya untuk apa pun yang mengikat.
* **``account.bank.statement.line`` tidak punya kolom ``date``** -- Odoo 19
  mendelegasikannya ke ``account.move``. Skrip ini memakai ORM, jadi aman; SQL
  mentah harus join ``account_move``.
* ``levis_trans_date`` adalah tanggal transaksi kartu menurut narasi, berbeda
  dari tanggal mutasi bank. Keduanya ditampilkan.
"""

import os
from collections import defaultdict
from datetime import date as date_cls

OUT = os.environ.get("OUT", "/mnt/data_levis/Laporan_MDR.xlsx")
DATE_FROM = os.environ.get("DATE_FROM", "2026-07-01")
DATE_TO = os.environ.get("DATE_TO", date_cls.today().isoformat())
MDR_ACCOUNT_CODE = os.environ.get("MDR_ACCOUNT", "7104000001")

company = env["res.company"].search([], order="id", limit=1)
currency = company.currency_id

print("=" * 78)
print("Laporan MDR dari mutasi bank   %s s/d %s" % (DATE_FROM, DATE_TO))
print("=" * 78)

lines = env["account.bank.statement.line"].search(
    [
        ("company_id", "=", company.id),
        ("date", ">=", DATE_FROM),
        ("date", "<=", DATE_TO),
        ("levis_mdr", "!=", 0),
    ],
    order="date, id",
)
print("\nBaris mutasi ber-MDR: %s" % len(lines))

per_month_bank = defaultdict(lambda: {"lines": 0, "gross": 0.0, "mdr": 0.0, "net": 0.0})
per_store_channel = defaultdict(lambda: {"lines": 0, "gross": 0.0, "mdr": 0.0})
no_store = []
detail = []
seen = defaultdict(list)

for line in lines:
    journal = line.journal_id
    store = line.levis_ou_analytic_id
    channel = line.levis_channel or "(tanpa tender)"
    month = line.date.strftime("%Y-%m")
    gross = line.levis_gross or 0.0
    mdr = line.levis_mdr or 0.0
    net = line.amount or 0.0

    bucket = per_month_bank[(month, journal.code or journal.display_name)]
    bucket["lines"] += 1
    bucket["gross"] += gross
    bucket["mdr"] += mdr
    bucket["net"] += net

    store_name = store.display_name if store else "(toko belum dipetakan)"
    sbucket = per_store_channel[(store_name, channel)]
    sbucket["lines"] += 1
    sbucket["gross"] += gross
    sbucket["mdr"] += mdr

    if not store:
        no_store.append(line)

    seen[(journal.id, line.date, round(net, 2), (line.payment_ref or "")[:60])].append(line.id)

    detail.append(
        [
            line.date,
            line.levis_trans_date or "",
            journal.code or journal.display_name,
            line.levis_mid or "",
            store_name,
            channel,
            gross,
            mdr,
            net,
            line.levis_narrative_kind or "",
            (line.payment_ref or "")[:120],
            "ya" if line.levis_clearing_line_id else "belum",
        ]
    )

print("\nPer bulan dan bank:")
print("%-9s %-8s %7s %18s %16s %18s" % ("Bulan", "Bank", "Baris", "Gross", "MDR", "Diterima"))
total_mdr = 0.0
for (month, bank), bucket in sorted(per_month_bank.items()):
    total_mdr += bucket["mdr"]
    print(
        "%-9s %-8s %7s %18s %16s %18s"
        % (
            month,
            bank,
            bucket["lines"],
            currency.round(bucket["gross"]),
            currency.round(bucket["mdr"]),
            currency.round(bucket["net"]),
        )
    )
print("%-18s %7s %18s %16s" % ("TOTAL", len(lines), "", currency.round(total_mdr)))

# --- sisi kedua: mutasi GL akun MDR ---
account = env["account.account"].with_company(company).search([("code", "=", MDR_ACCOUNT_CODE)], limit=1)
gl_by_month = defaultdict(float)
if account:
    gl_lines = env["account.move.line"].search(
        [
            ("company_id", "=", company.id),
            ("account_id", "=", account.id),
            ("parent_state", "=", "posted"),
            ("date", ">=", DATE_FROM),
            ("date", "<=", DATE_TO),
        ]
    )
    for gl in gl_lines:
        gl_by_month[gl.date.strftime("%Y-%m")] += gl.balance
    print("\nMutasi GL %s (%s): %s baris" % (MDR_ACCOUNT_CODE, account.display_name, len(gl_lines)))
else:
    print("\n*** Akun %s tidak ditemukan -- rekonsiliasi GL dilewati." % MDR_ACCOUNT_CODE)

narasi_by_month = defaultdict(float)
for (month, _bank), bucket in per_month_bank.items():
    narasi_by_month[month] += bucket["mdr"]

print("\nRekonsiliasi MDR narasi vs GL:")
print("%-9s %18s %18s %18s" % ("Bulan", "MDR narasi", "GL 7104000001", "Selisih"))
recon = []
for month in sorted(set(narasi_by_month) | set(gl_by_month)):
    narasi = narasi_by_month.get(month, 0.0)
    gl = gl_by_month.get(month, 0.0)
    recon.append([month, narasi, gl, narasi - gl])
    print("%-9s %18s %18s %18s" % (month, currency.round(narasi), currency.round(gl), currency.round(narasi - gl)))

duplicates = {key: ids for key, ids in seen.items() if len(ids) > 1}
print(
    "\nDugaan duplikat impor: %s kelompok, %s baris" % (len(duplicates), sum(len(ids) for ids in duplicates.values()))
)
print("Baris ber-MDR tanpa toko: %s" % len(no_store))

try:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Rincian"
    ws.append(
        [
            "Tgl Mutasi",
            "Tgl Transaksi",
            "Bank",
            "MID",
            "Toko",
            "Tender",
            "Gross",
            "MDR",
            "Diterima",
            "Jenis Narasi",
            "Narasi",
            "Sudah Clearing",
        ]
    )
    for row in detail:
        ws.append(row)

    w2 = wb.create_sheet("Per Bulan Bank")
    w2.append(["Bulan", "Bank", "Baris", "Gross", "MDR", "Diterima"])
    for (month, bank), bucket in sorted(per_month_bank.items()):
        w2.append([month, bank, bucket["lines"], bucket["gross"], bucket["mdr"], bucket["net"]])

    w3 = wb.create_sheet("Per Toko Tender")
    w3.append(["Toko", "Tender", "Baris", "Gross", "MDR"])
    for (store, channel), bucket in sorted(per_store_channel.items()):
        w3.append([store, channel, bucket["lines"], bucket["gross"], bucket["mdr"]])

    w4 = wb.create_sheet("Rekonsiliasi GL")
    w4.append(["Bulan", "MDR dari narasi bank", "Mutasi GL %s" % MDR_ACCOUNT_CODE, "Selisih"])
    for row in recon:
        w4.append(row)
    w4.append([])
    w4.append(
        [
            "Selisih bukan kesalahan laporan: ia berarti ada baris bank yang belum masuk "
            "clearing, atau clearing memakai angka lain."
        ]
    )

    w5 = wb.create_sheet("Tanpa Toko")
    w5.append(["Tgl Mutasi", "Bank", "MID", "Tender", "Gross", "MDR", "Narasi"])
    for line in no_store:
        w5.append(
            [
                line.date,
                line.journal_id.code,
                line.levis_mid or "",
                line.levis_channel or "",
                line.levis_gross or 0.0,
                line.levis_mdr or 0.0,
                (line.payment_ref or "")[:120],
            ]
        )
    if not no_store:
        w5.append(["(tidak ada)"])

    w6 = wb.create_sheet("Dugaan Duplikat")
    w6.append(["Bank", "Tanggal", "Jumlah", "Narasi", "Jumlah baris identik", "ID baris"])
    for (journal_id, when, amount, ref), ids in sorted(duplicates.items(), key=lambda kv: -len(kv[1])):
        w6.append(
            [
                env["account.journal"].browse(journal_id).code,
                when,
                amount,
                ref,
                len(ids),
                ", ".join(str(i) for i in ids),
            ]
        )
    if not duplicates:
        w6.append(["(tidak ada)"])

    wb.save(OUT)
    print("\nLaporan: %s" % OUT)
except Exception as exc:  # noqa: BLE001
    print("\n(laporan xlsx gagal: %s)" % str(exc)[:80])

env.cr.rollback()
