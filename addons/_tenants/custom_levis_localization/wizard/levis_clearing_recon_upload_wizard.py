# -*- coding: utf-8 -*-
"""``levis.clearing.recon.upload`` — the finished recon, coming back in.

Odoo maps what it can; the rest is a person's job, and that job happens in the
EBR workbook Finance already works in. This reads the ``UNMAPPED`` sheet of that
workbook back and turns the five input columns into durable decisions
(``levis.clearing.manual.map``), then recomputes the run so the decisions take
effect.

Three limits are deliberate:

* **It never books anything.** Apply ends at Compute. Generate Moves and Post
  stay what they have always been — a person who has read the summary.
* **It validates before it writes.** ``action_preview`` renders exactly what
  ``action_apply`` would do, and writes nothing at all, because a spreadsheet
  filled in over a week is the one input most likely to disagree with the ledger
  by the time it arrives.
* **A refusal is a sentence, not a silence.** A cash deposit pointed at a card
  receivable, a receipt already ticked elsewhere, a row whose amount has moved —
  each comes back naming the row and the reason, and is also persisted on
  ``levis.clearing.upload.log`` so it can be read again tomorrow.
"""

import base64
import hashlib
import io
import logging
from collections import defaultdict
from datetime import timedelta

from markupsafe import Markup

from odoo import _, api, fields, models
from odoo.exceptions import UserError
from odoo.tools import html_escape

_logger = logging.getLogger(__name__)

try:
    import openpyxl
except ImportError:  # pragma: no cover - the platform image ships openpyxl
    openpyxl = None

#: Column headers of the round-trip sheet, as written by ``levis.clearing.ebr``.
_KEY = "KEY"
_AMOUNT = "AMOUNT"
_DATE = "TANGGAL"
_STORE = "STORE CODE"
_RULE = "SIMPAN RULE (Y/N)"
_TENDER = "TENDER"
_REFS = "NO TRANSAKSI X24DN"
_NOTE = "CATATAN"


class LevisClearingReconUpload(models.TransientModel):
    _name = "levis.clearing.recon.upload"
    _description = "Upload Finished Clearing Recon"

    run_id = fields.Many2one("levis.pos.clearing", required=True, ondelete="cascade")
    company_id = fields.Many2one(related="run_id.company_id")
    file = fields.Binary(string="Recon Workbook", required=True, attachment=False)
    file_name = fields.Char(string="File Name")
    create_rules = fields.Boolean(
        string="Create MID rules",
        default=True,
        help="Where a row says Y and the bank line carries a MID or terminal, save "
        "the store choice as a mapping rule so next month resolves by itself.",
    )
    recompute = fields.Boolean(
        string="Recompute the run",
        default=True,
        help="Run Compute after applying, so the mappings take effect. Generate Moves and Post are untouched.",
    )
    force = fields.Boolean(
        string="Upload again anyway",
        help="Apply a file that has already been uploaded byte for byte.",
    )
    preview_html = fields.Html(string="Preview", readonly=True, sanitize=False)
    previewed = fields.Boolean(default=False)

    # ------------------------------------------------------------------
    # Buttons
    # ------------------------------------------------------------------
    def action_preview(self):
        """Read, validate, show. Writes nothing — not even a log row."""
        self.ensure_one()
        rows, meta = self._read_rows()
        results = self._validate(rows, meta)
        self.preview_html = self._render_preview(results, meta)
        self.previewed = True
        return self._reopen()

    def action_apply(self):
        self.ensure_one()
        rows, meta = self._read_rows()
        results = self._validate(rows, meta)
        log = self._write_log(results, meta)
        applied = [result for result in results if result["status"] == "apply"]
        rules = 0
        for result in applied:
            self._apply_row(result, log)
        if self.create_rules:
            rules = self._create_rules(applied)
        log.rule_count = rules
        if self.recompute and applied and self.run_id.state in ("draft", "computed"):
            self.run_id.action_compute()
            log.recomputed = True
        return {
            "type": "ir.actions.act_window",
            "res_model": "levis.clearing.upload.log",
            "res_id": log.id,
            "view_mode": "form",
            "target": "current",
        }

    def _reopen(self):
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }

    # ------------------------------------------------------------------
    # Reading the workbook
    # ------------------------------------------------------------------
    def _raw(self):
        self.ensure_one()
        if not self.file:
            raise UserError(_("Choose the recon workbook first."))
        return base64.b64decode(self.file)

    def _file_hash(self):
        return hashlib.sha256(self._raw()).hexdigest()

    def _read_rows(self):
        """``([row dicts], meta dict)`` from the uploaded workbook.

        Refusals here are about the *file*, not about its rows: a workbook from
        another run, or from a run that has since been generated, cannot be
        partially applied in any meaningful way.
        """
        self.ensure_one()
        if openpyxl is None:
            raise UserError(_("The openpyxl Python package is not installed on this server."))
        book = openpyxl.load_workbook(io.BytesIO(self._raw()), read_only=True, data_only=True)
        meta = self._read_meta(book)
        self._assert_file_applies(meta)
        if "UNMAPPED" not in book.sheetnames:
            raise UserError(_("This workbook has no UNMAPPED sheet — it is not the recon export Odoo produced."))
        sheet = book["UNMAPPED"]
        header_row = None
        headers = {}
        rows = []
        for index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = [("" if value is None else str(value).strip()) for value in values]
            if header_row is None:
                if cells and cells[0] == _KEY:
                    header_row = index
                    headers = {name: position for position, name in enumerate(cells) if name}
                    missing = [
                        name for name in (_KEY, _AMOUNT, _DATE, _STORE, _RULE, _TENDER, _REFS) if name not in headers
                    ]
                    if missing:
                        raise UserError(_("The UNMAPPED sheet is missing these columns: %s.", ", ".join(missing)))
                continue
            if not cells or not cells[headers[_KEY]]:
                continue
            rows.append(
                {
                    "row": index,
                    "key": cells[headers[_KEY]],
                    "amount": values[headers[_AMOUNT]],
                    "date": values[headers[_DATE]],
                    "store_code": cells[headers[_STORE]],
                    "save_rule": cells[headers[_RULE]].upper()[:1],
                    "tender": cells[headers[_TENDER]],
                    "refs": cells[headers[_REFS]],
                    "note": cells[headers[_NOTE]] if _NOTE in headers else "",
                }
            )
        book.close()
        if not rows:
            raise UserError(_("The UNMAPPED sheet has no rows to read."))
        return rows, meta

    def _read_meta(self, book):
        if "_META" not in book.sheetnames:
            raise UserError(_("This workbook has no _META sheet — only the export Odoo produced can be uploaded back."))
        meta = {}
        for values in book["_META"].iter_rows(values_only=True):
            if values and values[0]:
                meta[str(values[0]).strip()] = values[1]
        return meta

    def _assert_file_applies(self, meta):
        self.ensure_one()
        run = self.run_id
        if str(meta.get("run_id") or "") != str(run.id):
            raise UserError(
                _(
                    "This workbook was exported from run %(exported)s, but you are uploading it "
                    "into %(target)s. Open that run and upload it there.",
                    exported=meta.get("run_name") or meta.get("run_id") or "?",
                    target=run.name,
                )
            )
        if str(meta.get("company_id") or "") != str(run.company_id.id):
            raise UserError(_("This workbook belongs to another company."))
        if run.state not in ("draft", "computed"):
            raise UserError(
                _(
                    "%(run)s is already %(state)s. A mapping applied now would disagree with the "
                    "entries already planned — cancel it, or start a new run.",
                    run=run.name,
                    state=run.state,
                )
            )

    # ------------------------------------------------------------------
    # Validation — everything that can be refused, is refused here
    # ------------------------------------------------------------------
    def _validate(self, rows, meta):
        self.ensure_one()
        run = self.run_id
        company = run.company_id
        config = run.config_id or self.env["levis.clearing.config"]._get(company)
        cash_account = config._cash_receivable_account()
        tender_by_code = {
            (account.with_company(company).code or ""): account for account in config.pos_receivable_account_ids
        }
        lines_by_statement = {line.statement_line_id.id: line for line in run.line_ids}
        existing = self.env["levis.clearing.manual.map"]._for_lines(company, lines_by_statement.keys())
        Warehouse = self.env["stock.warehouse"]
        Receipt = self.env["levis.pos.clearing.receipt"]
        stale_token = str(meta.get("token") or "") != self.env["levis.clearing.ebr"]._token(run)
        claimed_here = {}
        results = []
        for row in rows:
            result = dict(row, status="skip", reason="", line=None, analytic=None, tender_account=None, ref_list=[])
            if not any((row["store_code"], row["tender"], row["refs"], row["note"])):
                result["reason"] = _("Nothing filled in.")
                results.append(result)
                continue
            line = lines_by_statement.get(self._as_int(row["key"]))
            if not line:
                result.update(status="reject", reason=_("This bank line is not part of %s.", run.name))
                results.append(result)
                continue
            result["line"] = line
            if not self._same_figures(line, row):
                result.update(
                    status="reject",
                    reason=_(
                        "The bank line now reads %(amount)s on %(date)s — the file was written "
                        "against different figures.",
                        amount="{:,.2f}".format(line.statement_amount),
                        date=line.settlement_date,
                    ),
                )
                results.append(result)
                continue
            if row["store_code"]:
                warehouse, analytic = Warehouse._levis_store_by_code(company, row["store_code"])
                if not warehouse:
                    result.update(status="reject", reason=_("Unknown store code %s.", row["store_code"]))
                    results.append(result)
                    continue
                if not analytic:
                    result.update(
                        status="reject",
                        reason=_("Store %s has no Operating Unit analytic account.", row["store_code"]),
                    )
                    results.append(result)
                    continue
                result["analytic"] = analytic
            if row["tender"]:
                account = tender_by_code.get(row["tender"].strip())
                if not account:
                    result.update(
                        status="reject",
                        reason=_("%s is not one of the configured POS tender receivables.", row["tender"]),
                    )
                    results.append(result)
                    continue
                refusal = self._tender_refusal(line, account, cash_account)
                if refusal:
                    result.update(status="reject", reason=refusal)
                    results.append(result)
                    continue
                result["tender_account"] = account
            if row["refs"]:
                refs = [ref.strip() for ref in row["refs"].replace(";", ",").split(",") if ref.strip()]
                refusal = self._refs_refusal(company, line, refs, Receipt, claimed_here)
                if refusal:
                    result.update(status="reject", reason=refusal)
                    results.append(result)
                    continue
                for ref in refs:
                    claimed_here[ref] = line
                result["ref_list"] = refs
            current = existing.get(line.statement_line_id.id)
            if current and self._unchanged(current, result):
                result.update(status="unchanged", reason=_("Already recorded."))
                results.append(result)
                continue
            result["status"] = "apply"
            if stale_token:
                result["reason"] = _("Applied — note the run was recomputed after this file was exported.")
            results.append(result)
        return results

    def _as_int(self, value):
        try:
            return int(float(str(value).strip()))
        except (TypeError, ValueError):
            return 0

    def _same_figures(self, line, row):
        """The row still describes the bank line it was exported from."""
        amount = row["amount"]
        try:
            if amount not in (None, "") and abs(float(amount) - line.statement_amount) > 0.005:
                return False
        except (TypeError, ValueError):
            return False
        wanted = row["date"]
        if wanted and line.settlement_date:
            wanted = getattr(wanted, "date", lambda: wanted)()
            if str(wanted)[:10] != str(line.settlement_date)[:10]:
                return False
        return True

    def _tender_refusal(self, line, account, cash_account):
        """The channel restriction, said out loud instead of applied silently.

        ``_pool_accounts_for_channel`` would simply ignore an account outside the
        channel's pool, and the person who typed it would never learn that their
        answer did nothing. So the same rule is checked here and refused by name.
        """
        if not cash_account:
            return ""
        if line.kind == "cash_deposit" and account != cash_account:
            return _(
                "%(entry)s is a cash deposit, so it settles %(cash)s and nothing else.",
                entry=line.move_name or "",
                cash=cash_account.display_name,
            )
        if line.kind != "cash_deposit" and account == cash_account:
            return _(
                "%(entry)s is card or QRIS money, so it may not settle the cash receivable.",
                entry=line.move_name or "",
            )
        return ""

    def _refs_refusal(self, company, line, refs, Receipt, claimed_here):
        known = self._known_refs(line)
        for ref in refs:
            if ref in claimed_here and claimed_here[ref] != line:
                return _("Transaction %s is claimed twice in this file.", ref)
            if known is not None and ref not in known:
                return _(
                    "Transaction %(ref)s is not among %(store)s's transactions for that trading day.",
                    ref=ref,
                    store=line.analytic_account_id.display_name or "",
                )
            claimed = Receipt.search(
                [
                    ("company_id", "=", company.id),
                    ("ref", "=", ref),
                    ("matched", "=", True),
                    ("line_id", "!=", line.id),
                ],
                limit=1,
            )
            if claimed and not claimed.suggested:
                # A human ticked it somewhere else. Two people disagreeing is not
                # something an upload may settle on its own.
                return _(
                    "Transaction %(ref)s is already matched by hand to %(entry)s.",
                    ref=ref,
                    entry=claimed.move_name or "",
                )
        return ""

    def _known_refs(self, line):
        """Every X24DN reference that store rang up around that trading day, or ``None``.

        ``None`` means the question cannot be asked — no store, no trading day, or
        no staged X70D at all — and the reference is then accepted on the
        operator's word rather than refused on data that is simply absent.

        The window is the same ladder the allocation may reach over, because a
        settlement is routinely paid a day or two after the trading day and a leg
        can legitimately name a receipt from either end of it.
        """
        analytic = line.analytic_account_id
        if not analytic or not line.trans_date:
            return None
        lookback = line.run_id.config_id.lookback_days or 10
        rows = self.env["levis.pos.clearing.alloc"]._x24_rows(
            {analytic.id},
            line.trans_date - timedelta(days=lookback),
            line.trans_date + timedelta(days=1),
            line.run_id.company_id,
        )
        if not rows:
            return None
        known = set()
        for (ou_id, _day), entries in rows.items():
            if ou_id != analytic.id:
                continue
            known.update(ref for _tender, ref, _amount in entries)
        return known or None

    def _unchanged(self, current, result):
        """The row says exactly what is already recorded, so applying it is a no-op.

        Only the fields the row actually filled in are compared: a second upload
        that carries the store but leaves the tender blank must not read as a
        change just because the first upload had set a tender.
        """
        if result["analytic"] and current.analytic_account_id != result["analytic"]:
            return False
        if result["tender_account"] and current.tender_account_id != result["tender_account"]:
            return False
        if result["ref_list"] and set(current._refs()) != set(result["ref_list"]):
            return False
        if result["note"] and (current.note or "") != result["note"]:
            return False
        return True

    # ------------------------------------------------------------------
    # Writing
    # ------------------------------------------------------------------
    def _apply_row(self, result, log):
        line = result["line"]
        ManualMap = self.env["levis.clearing.manual.map"]
        existing = ManualMap.search(
            [
                ("company_id", "=", self.company_id.id),
                ("statement_line_id", "=", line.statement_line_id.id),
            ],
            limit=1,
        )
        vals = {
            "company_id": self.company_id.id,
            "statement_line_id": line.statement_line_id.id,
            "statement_date": line.settlement_date,
            "statement_amount": line.statement_amount,
            "source": "upload",
            "upload_log_id": log.id,
            "user_id": self.env.user.id,
        }
        if result["analytic"]:
            vals["analytic_account_id"] = result["analytic"].id
        if result["tender_account"]:
            vals["tender_account_id"] = result["tender_account"].id
        if result["ref_list"]:
            vals["receipt_refs"] = ",".join(result["ref_list"])
        if result["note"]:
            vals["note"] = result["note"]
        if existing:
            existing.write(vals)
            return existing
        return ManualMap.create(vals)

    def _create_rules(self, applied):
        """Save the store choice as a mapping rule where the bank line names a key."""
        MidMap = self.env["levis.bank.mid.map"]
        wanted = []
        seen = set()
        for result in applied:
            if result["save_rule"] != "Y" or not result["analytic"]:
                continue
            line = result["line"]
            match_type, key = ("mid", line.mid_key) if line.mid_key else ("tid", line.tid_key)
            if not key:
                continue
            signature = (line.bank_journal_id.id, match_type, key)
            if signature in seen:
                continue
            seen.add(signature)
            if MidMap.search_count(
                [
                    ("company_id", "=", self.company_id.id),
                    ("journal_id", "=", line.bank_journal_id.id),
                    ("match_type", "=", match_type),
                    ("key", "=", key),
                ]
            ):
                continue
            wanted.append(
                {
                    "name": "%s %s — %s"
                    % (line.bank_journal_id.code or "", key, result["analytic"].display_name or ""),
                    "company_id": self.company_id.id,
                    "journal_id": line.bank_journal_id.id,
                    "match_type": match_type,
                    "key": key,
                    "channel": line.channel if line.channel in ("debit", "credit", "qris", "cash") else False,
                    "analytic_account_id": result["analytic"].id,
                    "note": _("Created from the recon upload of %s.", self.file_name or ""),
                }
            )
        created = 0
        for vals in wanted:
            # One at a time and inside a savepoint: a colliding rule must refuse
            # itself without taking the rest of the upload down with it.
            try:
                with self.env.cr.savepoint():
                    MidMap.create(vals)
                    created += 1
            except Exception as error:  # noqa: BLE001 - reported, not swallowed
                _logger.info("Recon upload: mapping rule %s refused: %s", vals.get("key"), error)
        return created

    def _write_log(self, results, meta):
        counts = defaultdict(int)
        for result in results:
            counts[result["status"]] += 1
        rejects = [result for result in results if result["status"] == "reject"]
        file_hash = self._file_hash()
        if not self.force:
            duplicate = self.env["levis.clearing.upload.log"].search(
                [("file_hash", "=", file_hash), ("state", "in", ("applied", "partial"))], limit=1
            )
            if duplicate:
                raise UserError(
                    _(
                        "This exact file was already uploaded on %(date)s (%(name)s). Tick "
                        '"Upload again anyway" if that is what you mean.',
                        date=duplicate.create_date,
                        name=duplicate.file_name,
                    )
                )
        state = "applied" if counts["apply"] and not rejects else ("partial" if counts["apply"] else "failed")
        log = self.env["levis.clearing.upload.log"].create(
            {
                "company_id": self.company_id.id,
                "run_id": self.run_id.id,
                "file_name": self.file_name or "recon.xlsx",
                "file_hash": file_hash,
                "state": state,
                "row_count": len(results),
                "applied_count": counts["apply"],
                "unchanged_count": counts["unchanged"],
                "rejected_count": len(rejects),
                "error_message": "\n".join(
                    "row %s (%s): %s" % (result["row"], result["key"], result["reason"]) for result in rejects
                )
                or False,
            }
        )
        self.env["ir.attachment"].create(
            {
                "name": self.file_name or "recon.xlsx",
                "res_model": log._name,
                "res_id": log.id,
                "type": "binary",
                "raw": self._raw(),
            }
        )
        return log

    # ------------------------------------------------------------------
    # Preview
    # ------------------------------------------------------------------
    def _render_preview(self, results, meta):
        counts = defaultdict(int)
        for result in results:
            counts[result["status"]] += 1
        head = _(
            "%(apply)s to apply, %(unchanged)s unchanged, %(reject)s refused, %(skip)s blank.",
            apply=counts["apply"],
            unchanged=counts["unchanged"],
            reject=counts["reject"],
            skip=counts["skip"],
        )
        parts = [Markup("<p><b>%s</b></p>") % head]
        shown = [result for result in results if result["status"] in ("apply", "reject", "unchanged")][:200]
        if shown:
            rows = Markup("").join(
                Markup("<tr><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td></tr>")
                % (
                    result["row"],
                    result["line"].move_name if result["line"] else result["key"],
                    result["status"],
                    result["store_code"] or (result["line"].analytic_account_id.display_name if result["line"] else ""),
                    result["reason"],
                )
                for result in shown
            )
            parts.append(
                Markup(
                    "<table class='table table-sm'><thead><tr>"
                    "<th>Row</th><th>Entry</th><th>Verdict</th><th>Store</th><th>Note</th>"
                    "</tr></thead><tbody>%s</tbody></table>"
                )
                % rows
            )
        if str(meta.get("token") or "") != self.env["levis.clearing.ebr"]._token(self.run_id):
            parts.append(
                Markup("<p class='text-warning'>%s</p>")
                % html_escape(
                    _(
                        "The run was recomputed after this file was exported. Every row is checked against the ledger as it stands now."
                    )
                )
            )
        return Markup("").join(parts)

    @api.onchange("file")
    def _onchange_file(self):
        self.previewed = False
        self.preview_html = False
