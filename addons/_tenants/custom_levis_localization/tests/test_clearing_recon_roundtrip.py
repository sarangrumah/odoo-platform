# -*- coding: utf-8 -*-
"""The recon workbook, out and back again.

The export is easy to test badly: assert that a file was produced and call it a
day. What actually matters is the round trip, and specifically the two things
that would make it worthless in production:

* a Compute after an upload must **keep** what the upload said — the clearing
  rebuilds its lines and its receipts from scratch every time, and that is
  exactly how 44 hand-made ticks were lost in September 2026;
* a refusal must be a sentence. ``_pool_accounts_for_channel`` would silently
  ignore a cash deposit pointed at a card receivable, and the person who typed it
  would never learn that their answer did nothing.

The fixture is the one from ``test_pos_clearing``, plus two warehouses, because
a store code is what the workbook is keyed on and only a warehouse carries one.
"""

import base64
import io
from datetime import date

import openpyxl

from odoo import Command
from odoo.addons.account.tests.common import AccountTestInvoicingCommon
from odoo.exceptions import UserError, ValidationError
from odoo.tests import tagged

MID_ONE = "885004600001"
MID_UNMAPPED = "885004600009"


@tagged("post_install", "-at_install")
class TestClearingReconRoundTrip(AccountTestInvoicingCommon):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.company_data["company"]
        Account = cls.env["account.account"]

        def account(name, code, kind, reconcile=False):
            return Account.create({"name": name, "code": code, "account_type": kind, "reconcile": reconcile})

        cls.suspense = account("Bank Suspense", "EBRSUS", "asset_current")
        cls.mdr = account("MDR Expense", "EBRMDR", "expense")
        cls.ar = account("Trade Receivable", "EBRAR", "asset_receivable", reconcile=True)
        cls.sweep = account("Main Bank", "EBRSWP", "asset_cash")
        cls.charge = account("Bank Charges", "EBRCHG", "expense")
        # The cash tender is resolved by code, so it has to carry the code the
        # parameter names — that is what makes the channel restriction real here.
        cls.tender_cash = account("POS Receivable - CASH", "1106000101", "asset_receivable", reconcile=True)
        cls.tender_visa = account("POS Receivable - OFFLINE_VISA", "1106000102", "asset_receivable", reconcile=True)
        cls.tenders = cls.tender_cash + cls.tender_visa

        plan = cls.env["account.analytic.plan"].create({"name": "EBR OU"})
        cls.store_one = cls.env["account.analytic.account"].create({"name": "STORE ONE", "plan_id": plan.id})
        cls.store_two = cls.env["account.analytic.account"].create({"name": "STORE TWO", "plan_id": plan.id})
        cls.warehouse_one = cls.env["stock.warehouse"].create(
            {
                "name": "Store One",
                "code": "EBR1",
                "company_id": cls.company.id,
                "l10n_store_code": "80001",
                "levis_ebr_label": "S-ONE",
                "l10n_ou_analytic_id": cls.store_one.id,
            }
        )
        cls.warehouse_two = cls.env["stock.warehouse"].create(
            {
                "name": "Store Two",
                "code": "EBR2",
                "company_id": cls.company.id,
                "l10n_store_code": "80002",
                "l10n_ou_analytic_id": cls.store_two.id,
            }
        )

        cls.gljv = cls.env["account.journal"].create({"name": "EBR Journal", "code": "EGLJ", "type": "general"})
        cls.bank = cls.env["account.journal"].create(
            {
                "name": "BCA test",
                "code": "EBCA",
                "type": "bank",
                "suspense_account_id": cls.suspense.id,
                "levis_clearing_format": "bca",
            }
        )
        cls.config = cls.env["levis.clearing.config"].create(
            {
                "company_id": cls.company.id,
                "journal_id": cls.gljv.id,
                "bank_journal_ids": [Command.set(cls.bank.ids)],
                "suspense_account_id": cls.suspense.id,
                "mdr_account_id": cls.mdr.id,
                "ar_account_id": cls.ar.id,
                "sweep_account_id": cls.sweep.id,
                "bank_charge_account_id": cls.charge.id,
                "pos_receivable_account_ids": [Command.set(cls.tenders.ids)],
                "settlement_lag_days": 1,
                "lookback_days": 10,
            }
        )
        cls.env["levis.bank.mid.map"].create(
            {
                "name": "Store one debit",
                "company_id": cls.company.id,
                "journal_id": cls.bank.id,
                "match_type": "mid",
                "key": MID_ONE,
                "channel": "debit",
                "analytic_account_id": cls.store_one.id,
            }
        )

    # ------------------------------------------------------------------
    # Fixture helpers
    # ------------------------------------------------------------------
    @classmethod
    def _posrec(cls, account, analytic, when, amount):
        move = cls.env["account.move"].create(
            {
                "move_type": "entry",
                "journal_id": cls.gljv.id,
                "company_id": cls.company.id,
                "date": when,
                "ref": "POS %s %s" % (analytic.name, when),
                "line_ids": [
                    Command.create(
                        {
                            "account_id": account.id,
                            "name": "POS receivable",
                            "debit": amount,
                            "credit": 0.0,
                            "analytic_distribution": {str(analytic.id): 100.0},
                        }
                    ),
                    Command.create(
                        {
                            "account_id": cls.company_data["default_account_revenue"].id,
                            "name": "Sales",
                            "debit": 0.0,
                            "credit": amount,
                        }
                    ),
                ],
            }
        )
        move.action_post()
        return move.line_ids.filtered(lambda aml: aml.account_id == account)

    @classmethod
    def _statement(cls, when, amount, payment_ref):
        line = cls.env["account.bank.statement.line"].create(
            {"journal_id": cls.bank.id, "date": when, "amount": amount, "payment_ref": payment_ref}
        )
        if line.move_id.state == "draft":
            line.move_id.action_post()
        return line

    @classmethod
    def _settlement_ref(cls, mid, gross, mdr):
        return "KR OTOMATIS MID : %s LEVIS TEST TGH: %.2f DDR: %.2f" % (mid, gross, mdr)

    def _run(self, **overrides):
        vals = {
            "company_id": self.company.id,
            "date_from": date(2026, 7, 1),
            "date_to": date(2026, 7, 31),
            "journal_id": self.gljv.id,
            "bank_journal_ids": [Command.set(self.bank.ids)],
        }
        vals.update(overrides)
        return self.env["levis.pos.clearing"].create(vals)

    # ------------------------------------------------------------------
    # Workbook helpers
    # ------------------------------------------------------------------
    def _export(self, run, **options):
        payload = {"compile_sales": True, "ar_sheet": True, "receipt_gaps": False}
        payload.update(options)
        return self.env["levis.clearing.ebr"]._build(run, options=payload)

    def _unmapped_rows(self, content):
        """``{statement line id: (row number, header index)}`` of the UNMAPPED sheet."""
        book = openpyxl.load_workbook(io.BytesIO(content))
        sheet = book["UNMAPPED"]
        header_row = None
        headers = {}
        rows = {}
        for index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = ["" if value is None else str(value).strip() for value in values]
            if header_row is None:
                if cells and cells[0] == "KEY":
                    header_row = index
                    headers = {name: position for position, name in enumerate(cells) if name}
                continue
            if cells and cells[0]:
                rows[int(float(cells[0]))] = index
        book.close()
        return rows, headers

    def _fill(self, content, edits):
        """Write ``{statement line id: {column header: value}}`` into UNMAPPED."""
        rows, headers = self._unmapped_rows(content)
        book = openpyxl.load_workbook(io.BytesIO(content))
        sheet = book["UNMAPPED"]
        for statement_id, values in edits.items():
            row = rows[statement_id]
            for header, value in values.items():
                sheet.cell(row=row, column=headers[header] + 1, value=value)
        out = io.BytesIO()
        book.save(out)
        return out.getvalue()

    def _upload(self, run, content, **overrides):
        vals = {"run_id": run.id, "file": base64.b64encode(content), "file_name": "recon.xlsx"}
        vals.update(overrides)
        return self.env["levis.clearing.recon.upload"].create(vals)

    # ------------------------------------------------------------------
    # The export
    # ------------------------------------------------------------------
    def test_export_carries_the_sheets_and_writes_nothing(self):
        self._posrec(self.tender_visa, self.store_one, date(2026, 7, 8), 1_000_000.0)
        self._statement(date(2026, 7, 9), 990_000.0, self._settlement_ref(MID_ONE, 1_000_000.0, 10_000.0))
        unmapped = self._statement(date(2026, 7, 9), 495_000.0, self._settlement_ref(MID_UNMAPPED, 500_000.0, 5_000.0))
        run = self._run()
        run.action_compute()

        before = (
            self.env["account.move"].search_count([("company_id", "=", self.company.id)]),
            len(run.line_ids),
            len(run.receipt_ids.filtered("matched")),
        )
        content = self._export(run)
        after = (
            self.env["account.move"].search_count([("company_id", "=", self.company.id)]),
            len(run.line_ids),
            len(run.receipt_ids.filtered("matched")),
        )
        self.assertEqual(before, after, "exporting the workbook must change nothing at all")

        book = openpyxl.load_workbook(io.BytesIO(content))
        self.assertIn("_META", book.sheetnames)
        self.assertIn("SUMMARY", book.sheetnames)
        self.assertIn("MUTASI BCA", book.sheetnames)
        self.assertIn("UNMAPPED", book.sheetnames)
        self.assertIn("COMPILE SALES", book.sheetnames)
        book.close()

        rows, _headers = self._unmapped_rows(content)
        self.assertIn(unmapped.id, rows, "the line nobody could map is the one the sheet is for")

    def test_mutasi_carries_what_odoo_resolved(self):
        self._posrec(self.tender_visa, self.store_one, date(2026, 7, 8), 1_000_000.0)
        self._statement(date(2026, 7, 9), 990_000.0, self._settlement_ref(MID_ONE, 1_000_000.0, 10_000.0))
        run = self._run()
        run.action_compute()

        book = openpyxl.load_workbook(io.BytesIO(self._export(run, compile_sales=False, ar_sheet=False)))
        sheet = book["MUTASI BCA"]
        header = None
        data = None
        for values in sheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value).strip() for value in values]
            if header is None:
                if cells and cells[0] == "Tanggal Transaksi":
                    header = {name: index for index, name in enumerate(cells) if name}
                continue
            if cells and cells[0] and not cells[0].startswith("TOTAL"):
                data = values
                break
        book.close()
        self.assertIsNotNone(data, "the bank sheet must carry its statement line")
        self.assertEqual(data[header["MID NO"]], MID_ONE)
        self.assertEqual(data[header["MDR"]], 10_000.0)
        self.assertEqual(data[header["AMOUNT PAYMENT"]], 1_000_000.0)
        self.assertEqual(data[header["Store code"]], "80001")
        self.assertEqual(data[header["STATUS"]], "REKON DONE")
        self.assertEqual(data[header["METHOD"]], "DEBIT")
        self.assertEqual(data[header["REMARKS"]], "S-ONE CEK (BCA)")

    def test_compile_sales_covers_the_trading_days_not_the_bank_days(self):
        """Money in on 1 September pays the day the store traded: 31 August.

        Filtering the sales side on the run's own dates reports the wrong days at
        both ends — it drops the takings the period actually settles and adds a
        day that will not be paid until the next period.
        """
        self._posrec(self.tender_visa, self.store_one, date(2026, 7, 8), 1_000_000.0)
        self._statement(date(2026, 7, 9), 990_000.0, self._settlement_ref(MID_ONE, 1_000_000.0, 10_000.0))
        run = self._run()
        run.action_compute()

        sales_from, sales_to = self.env["levis.clearing.ebr"]._trading_window(run)
        lag = run.config_id.settlement_lag_days
        self.assertEqual(lag, 1)
        self.assertEqual(sales_from, date(2026, 6, 30), "a July run settles trading days from 30 June")
        self.assertEqual(sales_to, date(2026, 7, 30))

        book = openpyxl.load_workbook(io.BytesIO(self._export(run)))
        banner = book["COMPILE SALES"].cell(row=1, column=1).value
        book.close()
        self.assertIn("2026-06-30", banner, "the sheet says which trading days it covers")
        self.assertIn("2026-07-30", banner)

    def test_a_tender_is_offered_by_name_and_read_back_either_way(self):
        """A bare code names nothing: ten receivables differ in their last digits."""
        Ebr = self.env["levis.clearing.ebr"]
        label = Ebr._tender_label(self.tender_visa, self.company)
        self.assertEqual(label, "1106000102 — POS Receivable - OFFLINE_VISA")

        day = date(2026, 7, 8)
        self._posrec(self.tender_visa, self.store_one, day, 300_000.0)
        self._posrec(self.tender_cash, self.store_one, day, 900_000.0)
        statement = self._statement(date(2026, 7, 9), 297_000.0, self._settlement_ref(MID_ONE, 300_000.0, 3_000.0))
        run = self._run()
        run.action_compute()

        # The label the sheet offers, and the bare code somebody typed instead:
        # both have to resolve to the same account.
        for written in (label, "1106000102"):
            self.env["levis.clearing.manual.map"].search([("statement_line_id", "=", statement.id)]).unlink()
            filled = self._fill(self._export(run, receipt_gaps=True), {statement.id: {"TENDER": written}})
            self._upload(run, filled, create_rules=False, force=True).action_apply()
            mapping = self.env["levis.clearing.manual.map"].search([("statement_line_id", "=", statement.id)])
            self.assertEqual(mapping.tender_account_id, self.tender_visa, "written as %r" % written)

    # ------------------------------------------------------------------
    # The round trip
    # ------------------------------------------------------------------
    def test_upload_maps_a_store_and_the_recompute_keeps_it(self):
        self._posrec(self.tender_visa, self.store_two, date(2026, 7, 8), 500_000.0)
        statement = self._statement(date(2026, 7, 9), 495_000.0, self._settlement_ref(MID_UNMAPPED, 500_000.0, 5_000.0))
        run = self._run()
        run.action_compute()
        self.assertEqual(run.line_ids.state, "unmapped")

        filled = self._fill(self._export(run), {statement.id: {"STORE CODE": "80002"}})
        wizard = self._upload(run, filled, create_rules=False)
        wizard.action_preview()
        self.assertIn("1 to apply", wizard.preview_html)
        wizard.action_apply()

        mapping = self.env["levis.clearing.manual.map"].search([("statement_line_id", "=", statement.id)])
        self.assertEqual(len(mapping), 1)
        self.assertEqual(mapping.analytic_account_id, self.store_two)
        self.assertEqual(mapping.source, "upload")

        line = run.line_ids
        self.assertEqual(line.analytic_account_id, self.store_two, "the upload already recomputed the run")
        self.assertEqual(line.state, "ok")
        self.assertEqual(line.manual_map_id, mapping)
        self.assertEqual(line.allocated, 500_000.0)

        # The point of the whole design: another Compute must not lose it.
        run.action_compute()
        self.assertEqual(run.line_ids.analytic_account_id, self.store_two)
        self.assertEqual(run.line_ids.state, "ok")

    def test_upload_locks_the_tender_it_names(self):
        day = date(2026, 7, 8)
        self._posrec(self.tender_visa, self.store_one, day, 300_000.0)
        self._posrec(self.tender_cash, self.store_one, day, 900_000.0)
        statement = self._statement(date(2026, 7, 9), 297_000.0, self._settlement_ref(MID_ONE, 300_000.0, 3_000.0))
        run = self._run()
        run.action_compute()

        # The line is mapped and allocated, so it is not on the default sheet:
        # that one is the decisions Odoo could not make. The switch is how a
        # tender gets overridden on a line that is merely wrong, not stuck.
        self.assertNotIn(statement.id, self._unmapped_rows(self._export(run))[0])
        filled = self._fill(self._export(run, receipt_gaps=True), {statement.id: {"TENDER": "1106000102"}})
        self._upload(run, filled, create_rules=False).action_apply()

        allocs = run.line_ids.alloc_ids
        self.assertEqual(allocs.account_id, self.tender_visa, "the named tender is drained first")
        self.assertTrue(run.line_ids.tender_locked)

    def test_manual_receipts_come_back_after_a_recompute(self):
        self._posrec(self.tender_visa, self.store_one, date(2026, 7, 8), 1_000_000.0)
        statement = self._statement(date(2026, 7, 9), 990_000.0, self._settlement_ref(MID_ONE, 1_000_000.0, 10_000.0))
        run = self._run()
        run.action_compute()

        filled = self._fill(
            self._export(run, receipt_gaps=True), {statement.id: {"NO TRANSAKSI X24DN": "80001-1-4711"}}
        )
        self._upload(run, filled, create_rules=False).action_apply()

        ticked = run.line_ids.receipt_ids.filtered("matched")
        self.assertEqual(ticked.mapped("ref"), ["80001-1-4711"])

        run.action_compute()
        ticked = run.line_ids.receipt_ids.filtered("matched")
        self.assertEqual(
            ticked.mapped("ref"),
            ["80001-1-4711"],
            "a recompute rebuilds the receipts, so the hand-made tick has to be put back",
        )

    def test_upload_can_save_the_choice_as_a_mapping_rule(self):
        self._posrec(self.tender_visa, self.store_two, date(2026, 7, 8), 500_000.0)
        statement = self._statement(date(2026, 7, 9), 495_000.0, self._settlement_ref(MID_UNMAPPED, 500_000.0, 5_000.0))
        run = self._run()
        run.action_compute()

        filled = self._fill(self._export(run), {statement.id: {"STORE CODE": "80002", "SIMPAN RULE (Y/N)": "Y"}})
        self._upload(run, filled).action_apply()

        rule = self.env["levis.bank.mid.map"].search([("company_id", "=", self.company.id), ("key", "=", MID_UNMAPPED)])
        self.assertEqual(len(rule), 1)
        self.assertEqual(rule.analytic_account_id, self.store_two)

    # ------------------------------------------------------------------
    # The refusals
    # ------------------------------------------------------------------
    def test_a_cash_deposit_may_not_be_pointed_at_a_card_receivable(self):
        statement = self._statement(date(2026, 7, 9), 400_000.0, "SETORAN TUNAI LEVIS TEST")
        run = self._run()
        run.action_compute()
        self.assertEqual(run.line_ids.kind, "cash_deposit")

        filled = self._fill(self._export(run), {statement.id: {"STORE CODE": "80001", "TENDER": "1106000102"}})
        wizard = self._upload(run, filled, create_rules=False)
        wizard.action_preview()
        self.assertIn("cash deposit", wizard.preview_html)
        wizard.action_apply()

        self.assertFalse(
            self.env["levis.clearing.manual.map"].search([("statement_line_id", "=", statement.id)]),
            "a refused row must not be half-applied",
        )
        log = self.env["levis.clearing.upload.log"].search([("run_id", "=", run.id)], limit=1)
        self.assertEqual(log.rejected_count, 1)
        self.assertEqual(log.state, "failed")

    def test_the_model_refuses_the_same_thing_directly(self):
        statement = self._statement(date(2026, 7, 9), 400_000.0, "SETORAN TUNAI LEVIS TEST")
        with self.assertRaises(ValidationError):
            self.env["levis.clearing.manual.map"].create(
                {
                    "company_id": self.company.id,
                    "statement_line_id": statement.id,
                    "analytic_account_id": self.store_one.id,
                    "tender_account_id": self.tender_visa.id,
                }
            )

    def test_a_row_whose_figures_moved_is_refused(self):
        statement = self._statement(date(2026, 7, 9), 495_000.0, self._settlement_ref(MID_UNMAPPED, 500_000.0, 5_000.0))
        run = self._run()
        run.action_compute()
        filled = self._fill(self._export(run), {statement.id: {"STORE CODE": "80002", "AMOUNT": 123_456.0}})
        wizard = self._upload(run, filled, create_rules=False)
        wizard.action_preview()
        self.assertIn("written against different figures", wizard.preview_html)

    def test_an_unknown_store_code_is_refused_by_name(self):
        statement = self._statement(date(2026, 7, 9), 495_000.0, self._settlement_ref(MID_UNMAPPED, 500_000.0, 5_000.0))
        run = self._run()
        run.action_compute()
        filled = self._fill(self._export(run), {statement.id: {"STORE CODE": "99999"}})
        wizard = self._upload(run, filled, create_rules=False)
        wizard.action_preview()
        self.assertIn("99999", wizard.preview_html)
        self.assertIn("Unknown store code", wizard.preview_html)

    def test_a_workbook_from_another_run_is_refused_whole(self):
        self._statement(date(2026, 7, 9), 495_000.0, self._settlement_ref(MID_UNMAPPED, 500_000.0, 5_000.0))
        run = self._run()
        run.action_compute()
        other = self._run(date_from=date(2026, 8, 1), date_to=date(2026, 8, 31))
        other.action_compute()
        with self.assertRaises(UserError):
            self._upload(other, self._export(run)).action_preview()

    def test_the_same_file_twice_is_refused_unless_forced(self):
        self._posrec(self.tender_visa, self.store_two, date(2026, 7, 8), 500_000.0)
        statement = self._statement(date(2026, 7, 9), 495_000.0, self._settlement_ref(MID_UNMAPPED, 500_000.0, 5_000.0))
        run = self._run()
        run.action_compute()
        filled = self._fill(self._export(run), {statement.id: {"STORE CODE": "80002"}})
        self._upload(run, filled, create_rules=False).action_apply()
        with self.assertRaises(UserError):
            self._upload(run, filled, create_rules=False).action_apply()
        # Forced, it goes through and reports the row as already recorded.
        wizard = self._upload(run, filled, create_rules=False, force=True)
        wizard.action_apply()
        log = self.env["levis.clearing.upload.log"].search([("run_id", "=", run.id)], order="id desc", limit=1)
        self.assertEqual(log.unchanged_count, 1)
        self.assertEqual(log.applied_count, 0)

    def test_a_generated_run_refuses_an_upload(self):
        self._posrec(self.tender_visa, self.store_one, date(2026, 7, 8), 1_000_000.0)
        self._statement(date(2026, 7, 9), 990_000.0, self._settlement_ref(MID_ONE, 1_000_000.0, 10_000.0))
        run = self._run()
        run.action_compute()
        content = self._export(run)
        run.action_generate_moves()
        with self.assertRaises(UserError):
            self._upload(run, content).action_preview()
